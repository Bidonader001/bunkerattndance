"""Payroll engine tests against a temporary SQLite database."""

from __future__ import annotations

import io
import os
import shutil
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


def _d(value) -> Decimal:
    return Decimal(str(value))


class PayrollTests(unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.mkdtemp(prefix="codai-pay-")
        os.environ["BUNKERATT_DB"] = str(Path(self._dir) / "test.db")
        from core import db, payroll

        self.db = db
        self.payroll = payroll
        db.init_db()
        with db.get_connection() as conn:
            conn.execute(
                "INSERT INTO gym_settings (key, value) VALUES ('grace_minutes', '10') ON CONFLICT(key) DO UPDATE SET value = excluded.value"
            )
            conn.execute(
                "INSERT INTO gym_settings (key, value) VALUES ('working_days', 'Sun-Thu') ON CONFLICT(key) DO UPDATE SET value = excluded.value"
            )
        db.add_employee("E1", "Amira Hassan", shift_id=1)
        db.add_employee("E2", "Omar Saleh", shift_id=1)
        payroll.save_policy(
            {
                "currency": "EGP",
                "divisor": "fixed_30",
                "hours_per_day": "8",
                "paid_leave_days": "4",
                "unused_leave_compensation": "0",
                "lateness_enabled": "1",
                "grace_minutes": "10",
                "late_method": "per_minute",
                "absence_full_day": "1",
                "absence_partial_enabled": "0",
                "early_departure_enabled": "0",
                "missing_checkout_policy": "ignore",
            },
            "tester",
        )
        payroll.save_profile("E1", {"monthly_salary": "10000", "effective_date": "2026-10-01"}, "tester")
        payroll.save_profile("E2", {"monthly_salary": "8000", "effective_date": "2026-10-01"}, "tester")

    def tearDown(self):
        os.environ.pop("BUNKERATT_DB", None)
        shutil.rmtree(self._dir, ignore_errors=True)

    def _session(self, eid, name, work_date, clock_in, clock_out=None, minutes_late=0, shift_name="Day"):
        hours = None
        if clock_out:
            hours = round(
                (datetime.fromisoformat(clock_out) - datetime.fromisoformat(clock_in)).total_seconds() / 3600.0,
                4,
            )
        with self.db.get_connection() as conn:
            conn.execute(
                """
                INSERT INTO work_sessions
                    (employee_id, employee_name, work_date, clock_in, clock_out, hours, confidence,
                     created_at, shift_name, late, minutes_late)
                VALUES (?, ?, ?, ?, ?, ?, 0.99, ?, ?, ?, ?)
                """,
                (
                    eid,
                    name,
                    work_date,
                    clock_in,
                    clock_out,
                    hours,
                    clock_in,
                    shift_name,
                    1 if minutes_late else 0,
                    minutes_late,
                ),
            )

    def _fill_present(self, eid="E1", name="Amira Hassan", start="2026-10-01", end="2026-10-31", skip=None):
        skip = set(skip or [])
        from datetime import timedelta

        cur = date.fromisoformat(start)
        last = date.fromisoformat(end)
        while cur <= last:
            key = cur.isoformat()
            if key not in skip:
                self._session(eid, name, key, key + "T08:00:00", key + "T16:00:00")
            cur = cur + timedelta(days=1)

    def _row(self, eid="E1", **policy):
        if policy:
            cur = self.payroll.get_policy()["values"]
            cur.update(policy)
            self.payroll.save_policy(cur, "tester")
        data = self.payroll.preview("2026-10-01", "2026-10-31", eid)
        for row in data["employees"]:
            if row["employee_id"] == eid:
                return row
        self.fail("employee missing from payroll")

    def test_10000_egp_fixed_30_divisor(self):
        self._fill_present()
        row = self._row()
        self.assertEqual(_d(row["base_salary"]), Decimal("10000.00"))
        self.assertEqual(_d(row["daily_rate"]), Decimal("333.33"))
        self.assertEqual(row["currency"], "EGP")

    def test_four_unused_leave_days_compensation_enabled(self):
        self._fill_present()
        row = self._row(unused_leave_compensation="1", unused_payout="end_of_period")
        self.assertEqual(_d(row["unused_leave_days"]), Decimal("4.00"))
        self.assertEqual(_d(row["unused_leave_compensation"]), Decimal("1333.32"))
        expected_net = Decimal("10000.00") + Decimal("1333.32")
        self.assertEqual(_d(row["net_pay"]), expected_net)

    def test_leave_compensation_disabled(self):
        self._fill_present()
        row = self._row(unused_leave_compensation="0")
        self.assertEqual(_d(row["unused_leave_compensation"]), Decimal("0.00"))
        self.assertEqual(_d(row["net_pay"]), Decimal("10000.00"))

    def test_partially_and_fully_used_leave(self):
        self._fill_present(skip={"2026-10-04", "2026-10-05", "2026-10-06", "2026-10-07"})
        self.payroll.add_leave("E1", "2026-10-04", "paid", "tester")
        self.payroll.add_leave("E1", "2026-10-05", "paid", "tester")
        row = self._row(unused_leave_compensation="1")
        self.assertEqual(_d(row["paid_leave_days"]), Decimal("2.00"))
        self.assertEqual(_d(row["unused_leave_days"]), Decimal("2.00"))
        self.payroll.add_leave("E1", "2026-10-06", "paid", "tester")
        self.payroll.add_leave("E1", "2026-10-07", "paid", "tester")
        row = self._row(unused_leave_compensation="1")
        self.assertEqual(_d(row["paid_leave_days"]), Decimal("4.00"))
        self.assertEqual(_d(row["unused_leave_days"]), Decimal("0.00"))
        self.assertEqual(_d(row["unused_leave_compensation"]), Decimal("0.00"))

    def test_paid_leave_is_not_unpaid_absence(self):
        self._fill_present(skip={"2026-10-04"})
        self.payroll.add_leave("E1", "2026-10-04", "paid", "tester")
        row = self._row()
        self.assertEqual(_d(row["paid_leave_days"]), Decimal("1.00"))
        self.assertEqual(_d(row["absence_days"]), Decimal("0.00"))
        self.assertEqual(_d(row["absence_deduction"]), Decimal("0.00"))

    def test_full_day_absence_deducts_daily_rate(self):
        self._fill_present(skip={"2026-10-08"})
        row = self._row()
        self.assertEqual(_d(row["absence_days"]), Decimal("1.00"))
        self.assertEqual(_d(row["absence_deduction"]), Decimal("333.33"))

    def test_lateness_day_fractions(self):
        daily = Decimal("333.33")
        cases = (
            (5, Decimal("0")),
            (15, Decimal("0")),
            (16, Decimal("0.25")),
            (30, Decimal("0.25")),
            (31, Decimal("0.50")),
            (60, Decimal("0.50")),
            (61, Decimal("1")),
        )
        days = ["2026-10-01", "2026-10-04", "2026-10-05", "2026-10-06", "2026-10-07", "2026-10-08", "2026-10-11"]
        self._fill_present(skip=set(days))
        fraction_sum = Decimal("0")
        chargeable = 0
        for work_date, (minutes, fraction) in zip(days, cases):
            self._session(
                "E1",
                "Amira Hassan",
                work_date,
                work_date + "T08:00:00",
                work_date + "T16:00:00",
                minutes_late=minutes,
            )
            fraction_sum += fraction
            chargeable += max(0, minutes - 15)
            row = self._row()
            expected = (daily * fraction_sum).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            self.assertEqual(_d(row["late_deduction"]), expected, minutes)
            self.assertEqual(row["chargeable_late_minutes"], chargeable, minutes)

    def test_monthly_late_cap(self):
        self._fill_present(skip={"2026-10-01"})
        self._session("E1", "Amira Hassan", "2026-10-01", "2026-10-01T10:00:00", "2026-10-01T16:00:00", minutes_late=120)
        row = self._row(late_monthly_cap="25")
        self.assertEqual(_d(row["late_deduction"]), Decimal("25.00"))

    def test_overnight_shift_session(self):
        self._fill_present(skip={"2026-10-04"})
        self._session(
            "E1",
            "Amira Hassan",
            "2026-10-04",
            "2026-10-04T20:00:00",
            "2026-10-05T04:00:00",
            shift_name="Night",
        )
        row = self._row()
        self.assertEqual(_d(row["absence_days"]), Decimal("0.00"))

    def test_missing_checkout_deducts_half_day(self):
        self._fill_present(skip={"2026-10-01"})
        self._session("E1", "Amira Hassan", "2026-10-01", "2026-10-01T08:00:00")
        row = self._row()
        self.assertEqual(_d(row["absence_deduction"]), Decimal("166.67"))
        self.assertEqual(_d(row["partial_absence_days"]), Decimal("0.50"))

    def test_team_leader_token_keeps_role(self):
        from core.auth import create_token, read_token

        self.assertEqual(read_token(create_token("sara", "leader"))["role"], "leader")

    def test_mark_absent_is_two_days_and_blocks_clock_in(self):
        self._fill_present(skip={"2026-10-08"})
        saved = self.payroll.mark_unpermitted_absence("E1", "2026-10-08", "leader")
        self.assertFalse(saved["already"])
        again = self.payroll.mark_unpermitted_absence("E1", "2026-10-08", "leader")
        self.assertTrue(again["already"])
        row = self._row()
        self.assertEqual(_d(row["absence_days"]), Decimal("2.00"))
        self.assertEqual(_d(row["absence_deduction"]), Decimal("666.66"))
        with self.assertRaises(self.payroll.PayrollError):
            self.payroll.mark_unpermitted_absence("E1", "2026-10-01", "leader")
        self.payroll.clear_unpermitted_absence("E1", "2026-10-08")
        cleared = self._row()
        self.assertEqual(_d(cleared["absence_deduction"]), Decimal("333.33"))

    def test_absent_without_permission_deducts_two_days(self):
        self._fill_present(skip={"2026-10-08"})
        self.payroll.add_leave("E1", "2026-10-08", "unpermitted", "tester")
        row = self._row()
        self.assertEqual(_d(row["absence_days"]), Decimal("2.00"))
        self.assertEqual(_d(row["absence_deduction"]), Decimal("666.66"))

    def test_late_and_missing_checkout_stack(self):
        self._fill_present(skip={"2026-10-01"})
        self._session("E1", "Amira Hassan", "2026-10-01", "2026-10-01T08:20:00", minutes_late=20)
        row = self._row()
        self.assertEqual(_d(row["late_deduction"]), Decimal("83.33"))
        self.assertEqual(_d(row["absence_deduction"]), Decimal("166.67"))

    def test_employee_override_salary(self):
        self._fill_present("E2", "Omar Saleh")
        row = self._row("E2")
        self.assertEqual(_d(row["base_salary"]), Decimal("8000.00"))
        self.assertTrue(row["inherited"]["leave_allowance"])

    def test_mid_period_salary_change(self):
        self._fill_present()
        self.payroll.save_profile("E1", {"monthly_salary": "16000", "effective_date": "2026-10-16"}, "tester")
        row = self._row()
        self.assertGreater(_d(row["base_salary"]), Decimal("10000.00"))
        self.assertLess(_d(row["base_salary"]), Decimal("16000.00"))

    def test_duplicate_calculate_and_finalized_lock(self):
        self._fill_present()
        first = self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")
        second = self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")
        self.assertEqual(first["period_id"], second["period_id"])
        self.payroll.set_status(first["period_id"], "finalized", "owner", "super")
        with self.assertRaises(self.payroll.PayrollError):
            self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")

    def test_duplicate_leave_payout_blocked(self):
        self._fill_present()
        self._row(unused_leave_compensation="1")
        period = self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")
        self.payroll.finalize(period["period_id"], "owner")
        again = self.payroll.preview("2026-10-01", "2026-10-31", "E1")
        row = again["employees"][0]
        self.assertEqual(_d(row["unused_leave_compensation"]), Decimal("0.00"))

    def test_carry_forward_cap(self):
        self._fill_present()
        self.payroll.save_policy(
            {**self.payroll.get_policy()["values"], "unused_leave_compensation": "0", "leave_carry_forward": "1", "max_carry_forward": "2"},
            "tester",
        )
        period = self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")
        self.payroll.finalize(period["period_id"], "owner")
        with self.db.get_connection() as conn:
            bal = conn.execute("SELECT carried FROM leave_balances WHERE employee_id = 'E1'").fetchone()
        self.assertEqual(_d(bal[0]), Decimal("2"))

    def test_rounding_matches_export_and_payslip(self):
        self._fill_present()
        period = self.payroll.calculate_period("2026-10-01", "2026-10-31", "tester")
        ui = period["employees"][0]
        slip = self.payroll.payslip(period["period_id"], "E1")
        payload, name, media = self.payroll.export_period(period["period_id"], "xlsx")
        wb = load_workbook(io.BytesIO(payload))
        ws = wb.active
        net_cell = None
        for row in ws.iter_rows(min_row=3, values_only=True):
            if row and row[0] == "E1":
                net_cell = str(row[17])
        self.assertEqual(ui["net_pay"], slip["net_pay"])
        self.assertEqual(ui["net_pay"], net_cell)
        self.assertTrue(name.endswith(".xlsx"))
        self.assertIn("spreadsheet", media)

    def test_part_time_charges_friday_and_pays_an_extra_monday(self):
        from datetime import timedelta

        with self.db.get_connection() as conn:
            conn.execute("UPDATE employees SET work_type = 'part' WHERE employee_id = 'E1'")
        cur = date(2026, 10, 1)
        while cur <= date(2026, 10, 31):
            if cur.weekday() in {3, 4, 5} and cur.isoformat() != "2026-10-02":
                key = cur.isoformat()
                self._session("E1", "Amira Hassan", key, key + "T08:00:00", key + "T16:00:00")
            cur = cur + timedelta(days=1)
        self._session("E1", "Amira Hassan", "2026-10-05", "2026-10-05T08:00:00", "2026-10-05T16:00:00")
        row = self._row()
        self.assertEqual(row["work_type"], "part")
        self.assertEqual(row["no_clock_in_days"], 1)
        self.assertEqual(_d(row["absence_deduction"]), Decimal("333.33"))
        self.assertEqual(row["extra_days"], 1)
        self.assertEqual(_d(row["extra_pay"]), Decimal("333.33"))
        board = self.payroll.month_board("2026-10-31")
        listed = next(item for item in board["employees"] if item["employee_id"] == "E1")
        self.assertEqual(listed["extra_days"], 1)
        self.assertEqual(_d(listed["to_pay"]), Decimal("10000.00"))

    def test_friday_without_clock_in_deducts_one_day(self):
        self._fill_present(skip={"2026-10-02"})
        row = self._row()
        self.assertEqual(row["no_clock_in_days"], 1)
        self.assertEqual(_d(row["absence_deduction"]), Decimal("333.33"))

    def test_day_off_allowance_and_fifth_day_refused(self):
        days = ["2026-10-02", "2026-10-03", "2026-10-09", "2026-10-10", "2026-10-16"]
        self._fill_present(skip=set(days))
        for work_date in days[:4]:
            saved = self.payroll.mark_day_off("E1", work_date, "admin")
            self.assertFalse(saved["already"])
        row = self._row()
        self.assertEqual(row["days_off_used"], 4)
        self.assertEqual(row["days_off_left"], 0)
        self.assertEqual(row["no_clock_in_days"], 1)
        self.assertEqual(_d(row["absence_deduction"]), Decimal("333.33"))
        with self.assertRaises(self.payroll.PayrollError) as caught:
            self.payroll.mark_day_off("E1", days[4], "admin")
        self.assertIn("No extra day off", str(caught.exception))

    def test_month_board_salary_and_counts(self):
        self._fill_present(skip={"2026-10-08"})
        self.payroll.mark_unpermitted_absence("E1", "2026-10-08", "admin")
        board = self.payroll.month_board("2026-10-08")
        row = next(item for item in board["employees"] if item["employee_id"] == "E1")
        self.assertEqual(row["salary"], "10000.00")
        self.assertEqual(row["absent"], 1)
        self.assertEqual(row["no_clock_in"], 0)
        self.assertEqual(row["days_off_left"], 4)
        self.assertEqual(_d(row["deducted"]), Decimal("666.66"))
        self.assertEqual(_d(row["to_pay"]), Decimal("9333.34"))

    def test_unauthorized_salary_api(self):
        from fastapi.testclient import TestClient

        from core.auth import create_token
        from server import app

        admin = create_token("tester", "admin")
        with TestClient(app) as client:
            denied = client.get("/api/v1/payroll/settings")
            self.assertEqual(denied.status_code, 401)
            ok = client.get("/api/v1/payroll/settings", headers={"Authorization": f"Bearer {admin}"})
            self.assertEqual(ok.status_code, 200)
            preview = client.get(
                "/api/v1/payroll/preview",
                params={"start_date": "2026-10-01", "end_date": "2026-10-31"},
                headers={"Authorization": f"Bearer {admin}"},
            )
            self.assertEqual(preview.status_code, 200)
            calc = client.post(
                "/api/v1/payroll/calculate",
                json={"start_date": "2026-10-01", "end_date": "2026-10-31"},
                headers={"Authorization": f"Bearer {admin}"},
            )
            self.assertEqual(calc.status_code, 200)
            pid = calc.json()["period_id"]
            locked = client.post(
                f"/api/v1/payroll/periods/{pid}/status",
                json={"status": "finalized"},
                headers={"Authorization": f"Bearer {admin}"},
            )
            self.assertEqual(locked.status_code, 403)


if __name__ == "__main__":
    unittest.main()

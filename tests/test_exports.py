"""Read-only Absence Less export tests against a temporary SQLite database."""

from __future__ import annotations

import io
import os
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from core.time_utils import hours_between


def _hours(clock_in: str, clock_out: str) -> float:
    start = datetime.fromisoformat(clock_in)
    end = datetime.fromisoformat(clock_out)
    return round(hours_between(start, end), 4)


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._dir = tempfile.mkdtemp(prefix="codai-export-")
        os.environ["BUNKERATT_DB"] = str(Path(cls._dir) / "test.db")
        from core import db

        db.init_db()
        with db.get_connection() as conn:
            conn.execute(
                "INSERT INTO gym_settings (key, value) VALUES ('grace_minutes', '10') ON CONFLICT(key) DO UPDATE SET value = excluded.value"
            )
            conn.execute(
                "INSERT INTO gym_settings (key, value) VALUES ('working_days', 'Sun-Thu') ON CONFLICT(key) DO UPDATE SET value = excluded.value"
            )
            conn.execute(
                "INSERT INTO gym_settings (key, value) VALUES ('company_name', 'Absence Less') ON CONFLICT(key) DO UPDATE SET value = excluded.value"
            )

        def person(eid, name, branch="HQ", department="Ops", status="active"):
            db.add_employee(eid, name)
            with db.get_connection() as conn:
                conn.execute(
                    "UPDATE employees SET branch = ?, department = ?, employment_status = ? WHERE employee_id = ?",
                    (branch, department, status, eid),
                )

        person("EMP-001", "Ahmed Mohamed")
        person("EMP-002", "Mariam Ali")
        person("EMP-003", "=2+2")
        person("EMP-004", "Inactive Person", status="inactive")
        person("EMP-005", "Other Branch", branch="Jeddah", department="Sales")
        person("EMP-006", "مريم علي")

        def session(eid, name, work_date, clock_in, clock_out=None, minutes_late=0, shift_name="Day"):
            hours = _hours(clock_in, clock_out) if clock_out else None
            with db.get_connection() as conn:
                conn.execute(
                    """
                    INSERT INTO work_sessions
                        (employee_id, employee_name, work_date, clock_in, clock_out, hours, confidence,
                         created_at, shift_name, late, minutes_late)
                    VALUES (?, ?, ?, ?, ?, ?, 0.99, ?, ?, ?, ?)
                    """,
                    (
                        eid,
                        name,
                        work_date,
                        clock_in,
                        clock_out,
                        hours,
                        clock_in,
                        shift_name,
                        1 if minutes_late else 0,
                        minutes_late,
                    ),
                )

        session("EMP-001", "Ahmed Mohamed", "2026-10-01", "2026-10-01T08:55:00", "2026-10-01T17:02:00")
        session("EMP-002", "Mariam Ali", "2026-10-01", "2026-10-01T09:18:00", "2026-10-01T17:10:00", minutes_late=18)
        session("EMP-003", "=2+2", "2026-10-01", "2026-10-01T08:00:00", "2026-10-01T16:00:00")
        session("EMP-006", "مريم علي", "2026-10-01", "2026-10-01T09:00:00", "2026-10-01T17:00:00")
        session("EMP-005", "Other Branch", "2026-10-01", "2026-10-01T08:00:00", "2026-10-01T16:00:00")
        session("EMP-001", "Ahmed Mohamed", "2026-10-02", "2026-10-02T10:00:00", "2026-10-02T14:00:00")
        session(
            "EMP-001",
            "Ahmed Mohamed",
            "2026-10-04",
            "2026-10-04T20:00:00",
            "2026-10-05T04:00:00",
            shift_name="Night",
        )
        session("EMP-001", "Ahmed Mohamed", "2026-10-05", "2026-10-05T08:00:00")
        session("EMP-001", "Ahmed Mohamed", "2026-10-06", "2026-10-06T08:00:00", "2026-10-06T12:00:00")
        session("EMP-001", "Ahmed Mohamed", "2026-10-06", "2026-10-06T13:00:00", "2026-10-06T17:00:00")

    @classmethod
    def tearDownClass(cls):
        os.environ.pop("BUNKERATT_DB", None)
        shutil.rmtree(cls._dir, ignore_errors=True)

    def _params(self, **extra):
        base = {
            "report": "monthly",
            "format": "xlsx",
            "start_date": "2026-10-01",
            "end_date": "2026-10-06",
            "branch": "HQ",
        }
        base.update(extra)
        return base

    def test_invalid_date_range(self):
        from core.exports import ExportError, validate_range

        with self.assertRaises(ExportError):
            validate_range("2026-10-31", "2026-10-01")
        with self.assertRaises(ExportError):
            validate_range("2026-01-01", "2027-12-31")

    def test_excel_workbook_sheets_and_columns(self):
        from core.exports import build

        payload, name, media = build(self._params(), "tester")
        self.assertTrue(name.endswith(".xlsx"))
        self.assertIn("spreadsheet", media)
        wb = load_workbook(io.BytesIO(payload))
        self.assertEqual(wb.sheetnames[0], "Attendance Details")
        self.assertIn("Employee Summary", wb.sheetnames)
        ws = wb["Attendance Details"]
        headers = [ws.cell(8, c).value for c in range(1, 8)]
        self.assertEqual(
            headers,
            ["Employee ID", "Name", "Date", "Check-in", "Check-out", "Hours", "Status"],
        )
        self.assertEqual(ws.freeze_panes, "A9")
        self.assertIsNotNone(ws.auto_filter.ref)

    def test_real_values_hours_and_late_grace(self):
        from core.exports import collect

        data = collect(self._params())
        by_key = {(r["employee_id"], r["date_iso"]): r for r in data["details"]}
        ahmed = by_key[("EMP-001", "2026-10-01")]
        mariam = by_key[("EMP-002", "2026-10-01")]
        self.assertEqual(ahmed["name"], "Ahmed Mohamed")
        self.assertEqual(ahmed["check_in"], "08:55")
        self.assertEqual(ahmed["check_out"], "17:02")
        self.assertEqual(ahmed["hours"], 8.12)
        self.assertEqual(ahmed["status"], "Present")
        self.assertEqual(mariam["status"], "Late")
        self.assertEqual(mariam["hours"], 7.87)

    def test_missing_checkout_stays_blank(self):
        from core.exports import collect

        data = collect(self._params())
        row = next(r for r in data["details"] if r["employee_id"] == "EMP-001" and r["date_iso"] == "2026-10-05")
        self.assertEqual(row["check_in"], "08:00")
        self.assertEqual(row["check_out"], "")
        self.assertIsNone(row["hours"])
        self.assertEqual(row["status"], "Incomplete")

    def test_overnight_shift_one_workday(self):
        from core.exports import collect

        data = collect(self._params())
        night = [r for r in data["details"] if r["employee_id"] == "EMP-001" and r["date_iso"] == "2026-10-04"]
        next_day = [r for r in data["details"] if r["employee_id"] == "EMP-001" and r["date_iso"] == "2026-10-05"]
        self.assertEqual(len(night), 1)
        self.assertEqual(night[0]["check_in"], "20:00")
        self.assertEqual(night[0]["check_out"], "04:00")
        self.assertEqual(night[0]["hours"], 8.00)
        self.assertEqual(len(next_day), 1)
        self.assertEqual(next_day[0]["status"], "Incomplete")

    def test_duplicate_scans_one_row(self):
        from core.exports import collect

        data = collect(self._params())
        rows = [r for r in data["details"] if r["employee_id"] == "EMP-001" and r["date_iso"] == "2026-10-06"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["check_in"], "08:00")
        self.assertEqual(rows[0]["check_out"], "17:00")
        self.assertEqual(rows[0]["hours"], 8.00)

    def test_rest_day_not_absence(self):
        from core.exports import collect

        data = collect(self._params(employee_id="EMP-002", end_date="2026-10-02"))
        summary = data["summary"][0]
        self.assertEqual(summary["days_present"], 1)
        self.assertEqual(summary["absences"], 0)
        self.assertTrue(any(r["date_iso"] == "2026-10-02" and r["employee_id"] == "EMP-001" for r in collect(self._params())["details"]))

    def test_summary_matches_details(self):
        from core.exports import collect

        data = collect(self._params(employee_id="EMP-001"))
        ahmed = next(r for r in data["details"] if r["employee_id"] == "EMP-001")
        self.assertTrue(ahmed)
        summary = data["summary"][0]
        present = {r["date_iso"] for r in data["details"] if r["status"] in {"Present", "Late", "Incomplete"}}
        scheduled_present = {d for d in present if d not in {"2026-10-02"}}
        self.assertEqual(summary["days_present"], len(scheduled_present))
        self.assertEqual(summary["late_arrivals"], 0)
        detail_hours = round(sum(r["hours"] or 0 for r in data["details"]), 2)
        self.assertEqual(summary["total_hours"], detail_hours)

    def test_inactive_excluded_unless_requested(self):
        from core.exports import collect

        data = collect(self._params(report="employees", branch=""))
        ids = {r["employee_id"] for r in data["directory"]}
        self.assertNotIn("EMP-004", ids)
        included = collect(self._params(report="employees", branch="", include_inactive="1"))
        self.assertIn("EMP-004", {r["employee_id"] for r in included["directory"]})

    def test_branch_restriction(self):
        from core.exports import collect

        hq = collect(self._params(report="employees", branch="HQ"))
        jed = collect(self._params(report="employees", branch="Jeddah"))
        self.assertNotIn("EMP-005", {r["employee_id"] for r in hq["directory"]})
        self.assertEqual([r["employee_id"] for r in jed["directory"]], ["EMP-005"])

    def test_formula_injection_csv_and_excel(self):
        from core.exports import build

        csv_payload, _, _ = build(self._params(format="csv"), "tester")
        self.assertTrue(csv_payload.startswith(b"\xef\xbb\xbf"))
        text = csv_payload.decode("utf-8-sig")
        self.assertIn("'=2+2", text)
        self.assertIn("مريم علي", text)
        xlsx, _, _ = build(self._params(), "tester")
        wb = load_workbook(io.BytesIO(xlsx))
        names = [row[1] for row in wb["Attendance Details"].iter_rows(min_row=9, max_col=2, values_only=True)]
        self.assertTrue(any(str(n).startswith("'=") for n in names if n))

    def test_empty_result_still_builds(self):
        from core.exports import build, collect

        data = collect(self._params(employee_id="EMP-NONE"))
        self.assertEqual(data["count"], 0)
        payload, name, _ = build(self._params(format="csv", employee_id="EMP-NONE"), "tester")
        self.assertTrue(name.endswith(".csv"))
        self.assertIn(b"Employee ID", payload)

    def test_pdf_and_absences_report(self):
        from core.exports import build, collect

        pdf, name, media = build(self._params(format="pdf"), "tester")
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(media, "application/pdf")
        self.assertIn("Summary", name)
        absences = collect(self._params(report="absences", start_date="2026-10-01", end_date="2026-10-05"))
        mariam_abs = [r for r in absences["details"] if r["employee_id"] == "EMP-002"]
        self.assertTrue(all(r["status"] == "Absent" for r in mariam_abs))
        self.assertFalse(any(r["date_iso"] == "2026-10-02" for r in mariam_abs))
        self.assertFalse(any(r["status"] == "On Leave" for r in absences["details"]))

    def test_export_does_not_mutate_sessions(self):
        from core import db
        from core.exports import build

        with db.get_connection() as conn:
            before = conn.execute("SELECT COUNT(*) FROM work_sessions").fetchone()[0]
        build(self._params(), "tester")
        with db.get_connection() as conn:
            after = conn.execute("SELECT COUNT(*) FROM work_sessions").fetchone()[0]
            audits = conn.execute("SELECT outcome FROM export_audit ORDER BY id DESC LIMIT 1").fetchone()[0]
        self.assertEqual(before, after)
        self.assertEqual(audits, "ok")

    def test_http_unauthorized_and_invalid_and_download(self):
        from fastapi.testclient import TestClient

        from core.auth import create_token
        from server import app

        token = create_token("tester", "admin")
        with TestClient(app) as client:
            denied = client.get("/api/v1/exports", params={"start_date": "2026-10-01", "end_date": "2026-10-06"})
            self.assertEqual(denied.status_code, 401)
            bad = client.get(
                "/api/v1/exports",
                params={"start_date": "2026-10-31", "end_date": "2026-10-01", "format": "xlsx"},
                headers={"Authorization": f"Bearer {token}"},
            )
            self.assertEqual(bad.status_code, 400)
            reports = client.get("/api/v1/reports/excel", params={"work_date": "2026-10-01", "period": "daily"})
            self.assertEqual(reports.status_code, 401)
            ok = client.get(
                "/api/v1/exports",
                params={
                    "report": "monthly",
                    "format": "xlsx",
                    "start_date": "2026-10-01",
                    "end_date": "2026-10-06",
                    "branch": "HQ",
                },
                headers={"Authorization": f"Bearer {token}"},
            )
            self.assertEqual(ok.status_code, 200)
            self.assertIn("attachment", ok.headers.get("content-disposition", ""))
            self.assertTrue(ok.content[:2] == b"PK")
            people = client.get(
                "/api/v1/exports/employees",
                params={"format": "csv", "branch": "HQ"},
                headers={"Authorization": f"Bearer {token}"},
            )
            self.assertEqual(people.status_code, 200)
            self.assertIn("EMP-001", people.text)
            self.assertNotIn("EMP-005", people.text)


if __name__ == "__main__":
    unittest.main()

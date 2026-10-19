"""Payroll engine: Decimal-safe salary, leave, and lateness rules. Read-only on attendance."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_CEILING, ROUND_DOWN, ROUND_HALF_UP, ROUND_UP
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from core import db
from core.exports import parse_working_days

TWOPLACES = Decimal("0.01")
ZERO = Decimal("0.00")

DIVISORS = ("fixed_30", "calendar_days", "scheduled_workdays")
LATE_METHODS = ("per_minute", "per_hour", "fixed", "progressive")
PAYOUTS = ("monthly", "end_of_period", "annual")
ROUNDING = {
    "half_up": ROUND_HALF_UP,
    "down": ROUND_DOWN,
    "up": ROUND_UP,
    "ceiling": ROUND_CEILING,
}
PERIOD_STATES = ("draft", "calculated", "pending_approval", "approved", "finalized")
LOCKED = frozenset({"finalized"})
LEAVE_KINDS = ("paid", "unpaid", "public_holiday", "weekly_rest", "exception", "late_ok", "unpermitted", "day_off")
LATE_GRACE_MINUTES = 15
DAYS_OFF_ALLOWANCE = 4
PART_TIME_WEEKDAYS = frozenset({3, 4, 5})  # Thursday, Friday, Saturday


def is_part_time(person: dict) -> bool:
    return str(person.get("work_type") or "full").strip().lower() == "part"


def is_scheduled(person: dict, day: date) -> bool:
    """Team leaders follow the full-time calendar. Part time is Thursday, Friday, and Saturday."""
    if not is_part_time(person):
        return True
    return day.weekday() in PART_TIME_WEEKDAYS

DEFAULT_POLICY = {
    "currency": "EGP",
    "divisor": "fixed_30",
    "hours_per_day": "8",
    "period": "monthly",
    "cutoff_day": "0",
    "rounding_precision": "2",
    "rounding": "half_up",
    "effective_date": "",
    "paid_leave_days": "0",
    "leave_period": "annual",
    "leave_accrual": "annual_reset",
    "unused_leave_compensation": "0",
    "unused_payout": "end_of_period",
    "leave_carry_forward": "0",
    "max_carry_forward": "0",
    "lateness_enabled": "1",
    "grace_minutes": "",
    "late_method": "per_minute",
    "late_fixed_amount": "0",
    "late_progressive": "[]",
    "late_daily_cap": "0",
    "late_monthly_cap": "0",
    "absence_full_day": "1",
    "absence_partial_enabled": "1",
    "early_departure_enabled": "1",
    "early_grace_minutes": "0",
    "missing_checkout_policy": "partial",
    "public_holidays": "",
}


class PayrollError(ValueError):
    pass


def _dec(value: Any, default: str = "0") -> Decimal:
    try:
        text = str(value if value is not None else default).strip() or default
        return Decimal(text)
    except Exception:
        return Decimal(default)


def _money(policy: dict, value: Decimal) -> Decimal:
    prec = max(0, min(4, _int(policy.get("rounding_precision"), 2)))
    quant = Decimal("1").scaleb(-prec)
    mode = ROUNDING.get(str(policy.get("rounding") or "half_up"), ROUND_HALF_UP)
    return value.quantize(quant, rounding=mode)


def _int(value: Any, default: int = 0) -> int:
    try:
        return int(Decimal(str(value if value is not None and str(value).strip() != "" else default)))
    except Exception:
        return default


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on"}


def _s(value: Any) -> str:
    return _money(DEFAULT_POLICY, _dec(value)).quantize(TWOPLACES, rounding=ROUND_HALF_UP).__format__("f")


def _now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def _dates(start: date, end: date) -> list[date]:
    out = []
    cur = start
    while cur <= end:
        out.append(cur)
        cur += timedelta(days=1)
    return out


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError as exc:
        raise PayrollError("Dates must be YYYY-MM-DD.") from exc


def _holidays(policy: dict) -> set[date]:
    out: set[date] = set()
    raw = str(policy.get("public_holidays") or "")
    for part in raw.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        try:
            out.add(date.fromisoformat(part[:10]))
        except ValueError:
            continue
    return out


def _company_grace() -> int:
    with db.get_connection() as conn:
        row = conn.execute("SELECT value FROM gym_settings WHERE key = 'grace_minutes'").fetchone()
    try:
        return max(0, int(str(row[0] if row else "10") or 10))
    except ValueError:
        return 10


def _working_days() -> set[int]:
    with db.get_connection() as conn:
        row = conn.execute("SELECT value FROM gym_settings WHERE key = 'working_days'").fetchone()
    return parse_working_days(str(row[0] if row else "Sun-Thu"))


def merge_policy(payload: dict | None = None) -> dict:
    out = dict(DEFAULT_POLICY)
    if payload:
        for key in DEFAULT_POLICY:
            if key in payload and payload[key] is not None:
                if key == "late_progressive" and not isinstance(payload[key], str):
                    out[key] = json.dumps(payload[key])
                else:
                    out[key] = str(payload[key]).strip()
    if not out.get("grace_minutes"):
        out["grace_minutes"] = str(_company_grace())
    if not out.get("effective_date"):
        out["effective_date"] = date.today().isoformat()
    if out.get("divisor") not in DIVISORS:
        out["divisor"] = "fixed_30"
    if out.get("late_method") not in LATE_METHODS:
        out["late_method"] = "per_minute"
    if out.get("unused_payout") not in PAYOUTS:
        out["unused_payout"] = "end_of_period"
    return out


def get_policy() -> dict:
    db.init_db()
    with db.get_connection() as conn:
        row = conn.execute("SELECT payload, version, effective_date, updated_at, updated_by FROM payroll_policy WHERE id = 1").fetchone()
    if not row:
        policy = merge_policy()
        return {"source": "default", "version": 0, "values": policy, "updated_at": "", "updated_by": ""}
    values = merge_policy(json.loads(row["payload"]))
    return {
        "source": "live",
        "version": int(row["version"] or 1),
        "values": values,
        "effective_date": row["effective_date"],
        "updated_at": row["updated_at"],
        "updated_by": row["updated_by"],
    }


def save_policy(payload: dict, actor: str) -> dict:
    db.init_db()
    current = get_policy()
    nxt = merge_policy({**current["values"], **(payload or {})})
    if nxt.get("divisor") not in DIVISORS:
        raise PayrollError("Divisor must be fixed_30, calendar_days, or scheduled_workdays.")
    blob = json.dumps(nxt, sort_keys=True)
    old = json.dumps(current["values"], sort_keys=True)
    stamp = _now()
    with db.get_connection() as conn:
        conn.execute(
            """
            INSERT INTO payroll_policy (id, payload, version, effective_date, updated_at, updated_by)
            VALUES (1, ?, 1, ?, ?, ?)
            ON CONFLICT(id) DO UPDATE SET
                payload = excluded.payload,
                version = payroll_policy.version + 1,
                effective_date = excluded.effective_date,
                updated_at = excluded.updated_at,
                updated_by = excluded.updated_by
            """,
            (blob, nxt["effective_date"], stamp, actor),
        )
        conn.execute(
            "INSERT INTO payroll_policy_audit (actor, old_value, new_value, effective_date, created_at) VALUES (?, ?, ?, ?, ?)",
            (actor, old, blob, nxt["effective_date"], stamp),
        )
        conn.execute(
            "INSERT INTO payroll_audit (actor, action, entity, entity_id, old_value, new_value, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (actor, "policy.update", "payroll_policy", "1", old, blob, stamp),
        )
    return get_policy()


def _audit(actor: str, action: str, entity: str, entity_id: str, old: str = "", new: str = "") -> None:
    with db.get_connection() as conn:
        conn.execute(
            "INSERT INTO payroll_audit (actor, action, entity, entity_id, old_value, new_value, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (actor, action, entity, entity_id, old[:8000], new[:8000], _now()),
        )


def get_profile(employee_id: str) -> dict:
    db.init_db()
    eid = (employee_id or "").strip()
    people = {p["employee_id"]: p for p in db.list_employees()}
    person = people.get(eid)
    if not person:
        raise PayrollError("Employee not found.")
    policy = get_policy()["values"]
    with db.get_connection() as conn:
        row = conn.execute("SELECT * FROM employee_payroll WHERE employee_id = ?", (eid,)).fetchone()
        hist = conn.execute(
            "SELECT monthly_salary, effective_date FROM employee_salary_history WHERE employee_id = ? ORDER BY effective_date",
            (eid,),
        ).fetchall()
    inherited = {
        "monthly_salary": True,
        "divisor": True,
        "leave_allowance": True,
        "shift": True,
        "unused_leave_eligible": True,
    }
    monthly = ZERO
    divisor = None
    leave_allow = None
    shift_id = person.get("shift_id")
    unused_el = None
    allowances = []
    effective = person.get("joining_date") or ""
    if row:
        monthly = _dec(row["monthly_salary"])
        inherited["monthly_salary"] = monthly == ZERO
        if row["divisor_override"]:
            divisor = row["divisor_override"]
            inherited["divisor"] = False
        if row["leave_allowance_override"] not in (None, ""):
            leave_allow = _dec(row["leave_allowance_override"])
            inherited["leave_allowance"] = False
        if row["shift_id_override"] is not None:
            shift_id = row["shift_id_override"]
            inherited["shift"] = False
        if row["unused_leave_eligible"] is not None:
            unused_el = bool(row["unused_leave_eligible"])
            inherited["unused_leave_eligible"] = False
        try:
            allowances = json.loads(row["allowances_json"] or "[]")
        except json.JSONDecodeError:
            allowances = []
        effective = row["effective_date"] or effective
    return {
        "employee_id": eid,
        "name": person.get("name") or eid,
        "monthly_salary": str(_money(policy, monthly)),
        "divisor_override": divisor,
        "leave_allowance_override": None if leave_allow is None else str(leave_allow),
        "shift_id": shift_id,
        "shift_name": person.get("shift_name") or "",
        "unused_leave_eligible": unused_el,
        "effective_date": effective,
        "allowances": allowances if isinstance(allowances, list) else [],
        "salary_history": [dict(h) for h in hist],
        "inherited": inherited,
        "joining_date": person.get("joining_date") or "",
        "employment_status": person.get("employment_status") or "active",
    }


def save_profile(employee_id: str, payload: dict, actor: str) -> dict:
    db.init_db()
    eid = (employee_id or "").strip()
    current = get_profile(eid)
    salary = _dec(payload.get("monthly_salary"), current["monthly_salary"])
    if salary < 0:
        raise PayrollError("Salary cannot be negative.")
    divisor = payload.get("divisor_override")
    if divisor in ("", None, "inherit"):
        divisor = None
    if divisor is not None and divisor not in DIVISORS:
        raise PayrollError("Invalid salary divisor override.")
    leave = payload.get("leave_allowance_override")
    leave_s = None if leave in ("", None, "inherit") else str(_dec(leave))
    shift_raw = payload.get("shift_id_override", payload.get("shift_id"))
    shift_id = None if shift_raw in ("", None, "inherit") else _int(shift_raw, 0) or None
    unused = payload.get("unused_leave_eligible")
    unused_i = None if unused in ("", None, "inherit") else (1 if _bool(unused) else 0)
    effective = str(payload.get("effective_date") or date.today().isoformat())[:10]
    allowances = payload.get("allowances") if isinstance(payload.get("allowances"), list) else current["allowances"]
    stamp = _now()
    with db.get_connection() as conn:
        conn.execute(
            """
            INSERT INTO employee_payroll
                (employee_id, monthly_salary, divisor_override, leave_allowance_override, shift_id_override,
                 unused_leave_eligible, effective_date, allowances_json, updated_at, updated_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(employee_id) DO UPDATE SET
                monthly_salary = excluded.monthly_salary,
                divisor_override = excluded.divisor_override,
                leave_allowance_override = excluded.leave_allowance_override,
                shift_id_override = excluded.shift_id_override,
                unused_leave_eligible = excluded.unused_leave_eligible,
                effective_date = excluded.effective_date,
                allowances_json = excluded.allowances_json,
                updated_at = excluded.updated_at,
                updated_by = excluded.updated_by
            """,
            (eid, str(salary), divisor, leave_s, shift_id, unused_i, effective, json.dumps(allowances), stamp, actor),
        )
        if str(salary) != str(_dec(current["monthly_salary"])):
            conn.execute(
                "INSERT INTO employee_salary_history (employee_id, monthly_salary, effective_date, actor, created_at) VALUES (?, ?, ?, ?, ?)",
                (eid, str(salary), effective, actor, stamp),
            )
    _audit(actor, "profile.update", "employee_payroll", eid, json.dumps(current)[:4000], json.dumps(payload)[:4000])
    return get_profile(eid)


def mark_unpermitted_absence(employee_id: str, work_date: str, actor: str) -> dict:
    eid = (employee_id or "").strip()
    person = next((p for p in db.list_employees() if p["employee_id"] == eid), None)
    if person is None:
        raise PayrollError("Employee not found.")
    if (person.get("employment_status") or "active").lower() == "inactive":
        raise PayrollError("Inactive employees cannot be marked absent.")
    day = _parse_date(work_date).isoformat()
    if db.list_sessions(work_date=day, employee_id=eid):
        raise PayrollError("This employee already clocked in on that day.")
    if any(row["employee_id"] == eid for row in db.list_day_leave(day, "day_off")):
        raise PayrollError("This day is already a day off.")
    if any(row["employee_id"] == eid for row in db.list_day_leave(day, "unpermitted")):
        return {"ok": True, "employee_id": eid, "name": person.get("name") or eid, "work_date": day, "already": True}
    saved = add_leave(eid, day, "unpermitted", actor, note="Marked absent")
    saved["name"] = person.get("name") or eid
    saved["already"] = False
    return saved


def clear_unpermitted_absence(employee_id: str, work_date: str) -> dict:
    eid = (employee_id or "").strip()
    day = _parse_date(work_date).isoformat()
    removed = db.delete_day_leave(eid, day, "unpermitted")
    return {"ok": True, "removed": removed, "employee_id": eid, "work_date": day}


def _month_span(day: date) -> tuple[date, date]:
    start = day.replace(day=1)
    last = (start + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    return start, last


def days_off_used(employee_id: str, work_date: str) -> int:
    day = _parse_date(work_date)
    start, last = _month_span(day)
    with db.get_connection() as conn:
        row = conn.execute(
            """
            SELECT COUNT(*) FROM leave_ledger
            WHERE employee_id = ? AND kind = 'day_off' AND work_date >= ? AND work_date <= ?
            """,
            ((employee_id or "").strip(), start.isoformat(), last.isoformat()),
        ).fetchone()
    return int(row[0] if row else 0)


def days_off_left(employee_id: str, work_date: str) -> int:
    return max(0, DAYS_OFF_ALLOWANCE - days_off_used(employee_id, work_date))


def mark_day_off(employee_id: str, work_date: str, actor: str) -> dict:
    eid = (employee_id or "").strip()
    person = next((p for p in db.list_employees() if p["employee_id"] == eid), None)
    if person is None:
        raise PayrollError("Employee not found.")
    if (person.get("employment_status") or "active").lower() == "inactive":
        raise PayrollError("Inactive employees cannot take a day off.")
    day = _parse_date(work_date).isoformat()
    if db.list_sessions(work_date=day, employee_id=eid):
        raise PayrollError("This employee already clocked in on that day.")
    if any(row["employee_id"] == eid for row in db.list_day_leave(day, "unpermitted")):
        raise PayrollError("This day is already marked absent.")
    if any(row["employee_id"] == eid for row in db.list_day_leave(day, "day_off")):
        saved = {"ok": True, "employee_id": eid, "name": person.get("name") or eid, "work_date": day, "already": True}
        saved["days_off_left"] = days_off_left(eid, day)
        return saved
    if days_off_used(eid, day) >= DAYS_OFF_ALLOWANCE:
        raise PayrollError("No extra day off.")
    saved = add_leave(eid, day, "day_off", actor, note="Day off")
    saved["name"] = person.get("name") or eid
    saved["already"] = False
    saved["days_off_left"] = days_off_left(eid, day)
    return saved


def clear_day_off(employee_id: str, work_date: str) -> dict:
    eid = (employee_id or "").strip()
    day = _parse_date(work_date).isoformat()
    removed = db.delete_day_leave(eid, day, "day_off")
    return {"ok": True, "removed": removed, "employee_id": eid, "work_date": day, "days_off_left": days_off_left(eid, day)}


def add_leave(employee_id: str, work_date: str, kind: str, actor: str, note: str = "", days: str = "1") -> dict:
    db.init_db()
    eid = (employee_id or "").strip()
    if kind not in LEAVE_KINDS:
        raise PayrollError("Unknown leave kind.")
    day = _parse_date(work_date).isoformat()
    stamp = _now()
    with db.get_connection() as conn:
        conn.execute(
            "INSERT INTO leave_ledger (employee_id, work_date, kind, days, minutes, note, actor, created_at) VALUES (?, ?, ?, ?, 0, ?, ?, ?)",
            (eid, day, kind, str(_dec(days, "1")), note, actor, stamp),
        )
        lid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    _audit(actor, "leave.add", "leave_ledger", str(lid), "", f"{eid} {day} {kind}")
    return {"id": lid, "employee_id": eid, "work_date": day, "kind": kind}


def add_adjustment(employee_id: str, amount: str, kind: str, actor: str, note: str = "", period_id: int | None = None) -> dict:
    db.init_db()
    if period_id:
        period = _period(period_id)
        if period["status"] in LOCKED:
            raise PayrollError("Finalized payroll cannot take silent adjustments. Open a new correction period.")
    amt = _dec(amount)
    k = (kind or "bonus").strip().lower()
    if k not in {"bonus", "allowance", "deduction", "correction"}:
        raise PayrollError("Adjustment kind must be bonus, allowance, deduction, or correction.")
    stamp = _now()
    with db.get_connection() as conn:
        conn.execute(
            "INSERT INTO payroll_adjustments (period_id, employee_id, kind, amount, note, actor, created_at, status) VALUES (?, ?, ?, ?, ?, ?, ?, 'approved')",
            (period_id, employee_id.strip(), k, str(amt), note, actor, stamp),
        )
        aid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    _audit(actor, "adjustment.add", "payroll_adjustments", str(aid), "", str(amt))
    return {"id": aid, "employee_id": employee_id, "kind": k, "amount": str(amt)}


def _period(period_id: int) -> dict:
    with db.get_connection() as conn:
        row = conn.execute("SELECT * FROM payroll_periods WHERE id = ?", (period_id,)).fetchone()
    if not row:
        raise PayrollError("Payroll period not found.")
    return dict(row)


def _salary_on(history: list[dict], fallback: Decimal, day: date) -> Decimal:
    chosen = fallback
    for row in history:
        try:
            when = date.fromisoformat(str(row.get("effective_date") or "")[:10])
        except ValueError:
            continue
        if when <= day:
            chosen = _dec(row.get("monthly_salary"), str(fallback))
    return chosen


def _divisor_n(policy: dict, start: date, end: date, working: set[int], holidays: set[date]) -> Decimal:
    kind = policy.get("divisor") or "fixed_30"
    days = _dates(start, end)
    if kind == "calendar_days":
        return Decimal(max(1, len(days)))
    if kind == "scheduled_workdays":
        n = sum(1 for d in days if d.weekday() in working and d not in holidays)
        return Decimal(max(1, n))
    return Decimal("30")


def _late_day_fraction(minutes_late: int) -> Decimal:
    """Day fraction deducted for lateness. 6:15 on a 6:00 shift is still on time."""
    minutes = max(0, int(minutes_late))
    if minutes <= LATE_GRACE_MINUTES:
        return ZERO
    if minutes <= 30:
        return Decimal("0.25")
    if minutes <= 60:
        return Decimal("0.50")
    return Decimal("1")


def calculate_employee(
    person: dict,
    policy: dict,
    start: date,
    end: date,
    sessions: list[dict],
    leave_rows: list[dict],
    adjustments: list[dict],
    history: list[dict],
    working: set[int],
    holidays: set[date],
    shifts: list[dict],
    paid_out_days: Decimal,
    carried: Decimal,
    *,
    payout_now: bool,
) -> dict:
    profile_salary = _dec(person.get("monthly_salary"))
    hours_day = max(Decimal("1"), _dec(policy.get("hours_per_day"), "8"))
    divisor_n = _divisor_n(policy, start, end, working, holidays)
    allowance = _dec(person.get("leave_allowance"), policy.get("paid_leave_days"))
    if person.get("leave_allowance_override") not in (None, ""):
        allowance = _dec(person.get("leave_allowance_override"))
    allowance = allowance + carried
    unused_ok = person.get("unused_leave_eligible")
    if unused_ok is None:
        unused_ok = _bool(policy.get("unused_leave_compensation"))
    else:
        unused_ok = bool(unused_ok)

    by_day: dict[str, list[dict]] = {}
    for sess in sessions:
        by_day.setdefault(str(sess.get("work_date") or "")[:10], []).append(sess)
    leave_by: dict[str, list[dict]] = {}
    for row in leave_rows:
        leave_by.setdefault(str(row.get("work_date") or "")[:10], []).append(row)

    join = None
    if person.get("joining_date"):
        try:
            join = date.fromisoformat(str(person["joining_date"])[:10])
        except ValueError:
            join = None

    paid_leave = ZERO
    unpaid_leave = ZERO
    absence_days = ZERO
    partial_days = ZERO
    late_minutes = 0
    late_days = 0
    absent_marks = 0
    no_clock_in_days = 0
    no_clock_out_days = 0
    days_off_used_n = 0
    extra_days = 0
    extra_pay = ZERO
    chargeable_total = 0
    late_ded = ZERO
    early_ded = ZERO
    absence_ded = ZERO
    unpaid_ded = ZERO
    missing_flags: list[str] = []
    explanation: list[str] = []
    lines: list[dict] = []
    day_salary_sum = ZERO
    calendar_days = _dates(start, end)

    def _charge_presence(group: list[dict], kinds: set[str], daily: Decimal, key: str) -> None:
        nonlocal late_minutes, late_days, chargeable_total, late_ded
        nonlocal no_clock_out_days, absence_ded, partial_days
        late_min = max(int(s.get("minutes_late") or 0) for s in group)
        late_minutes += late_min
        approved_late = "exception" in kinds or "late_ok" in kinds
        if _bool(policy.get("lateness_enabled")) and not approved_late:
            fraction = _late_day_fraction(late_min)
            chargeable_total += max(0, late_min - LATE_GRACE_MINUTES)
            if fraction > 0:
                late_days += 1
                amt = daily * fraction
                cap = _dec(policy.get("late_daily_cap"))
                if cap > 0:
                    amt = min(amt, cap)
                late_ded += amt
                explanation.append(f"{key}: {late_min} minutes late — {fraction} day deducted.")
        elif approved_late:
            explanation.append(f"{key}: late arrival approved — no lateness deduction.")
        open_rows = [s for s in group if not s.get("clock_out")]
        if open_rows:
            missing_flags.append(key)
            no_clock_out_days += 1
            absence_ded += daily * Decimal("0.5")
            partial_days += Decimal("0.5")
            explanation.append(f"{key}: no clock-out — 0.5 day deducted.")

    for day in calendar_days:
        key = day.isoformat()
        sal = _salary_on(history, profile_salary, day)
        daily = _money(policy, sal / divisor_n) if divisor_n else ZERO
        day_salary_sum += sal
        if join and day < join:
            continue
        kinds = {str(r.get("kind")) for r in leave_by.get(key, [])}
        if day in holidays or "public_holiday" in kinds:
            continue
        if "weekly_rest" in kinds:
            continue
        group = by_day.get(key) or []
        if not is_scheduled(person, day):
            if not group:
                continue
            extra_days += 1
            extra_pay += daily
            explanation.append(f"{key}: extra day — 1 day added.")
            _charge_presence(group, kinds, daily, key)
            continue
        if "paid" in kinds:
            paid_leave += _dec(leave_by[key][0].get("days"), "1")
            explanation.append(f"{key}: paid leave — not an unpaid absence.")
            continue
        if "unpaid" in kinds:
            days_n = _dec(leave_by[key][0].get("days"), "1")
            unpaid_leave += days_n
            unpaid_ded += daily * days_n
            continue
        if "unpermitted" in kinds:
            days_n = _dec(leave_by[key][0].get("days"), "1")
            if days_n <= 0:
                days_n = Decimal("1")
            penalty_days = Decimal("2") * days_n
            absence_days += penalty_days
            absent_marks += 1
            absence_ded += daily * penalty_days
            explanation.append(f"{key}: absent without permission — {penalty_days} days deducted.")
            continue
        if "day_off" in kinds and days_off_used_n < DAYS_OFF_ALLOWANCE:
            days_off_used_n += 1
            explanation.append(f"{key}: day off — no deduction.")
            continue
        group = by_day.get(key) or []
        if not group:
            if _bool(policy.get("absence_full_day")):
                absence_days += Decimal("1")
                no_clock_in_days += 1
                absence_ded += daily
                if "day_off" in kinds:
                    explanation.append(f"{key}: no extra day off — 1 day deducted.")
                else:
                    explanation.append(f"{key}: no clock-in — 1 day deducted.")
            continue
        _charge_presence(group, kinds, daily, key)

    month_cap = _dec(policy.get("late_monthly_cap"))
    if month_cap > 0:
        late_ded = min(late_ded, month_cap)

    n_cal = Decimal(max(1, len(calendar_days)))
    avg_salary = day_salary_sum / n_cal
    last_dom = (date(start.year, start.month, 1) + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    full_month = start.day == 1 and end == last_dom
    daily_rate = _money(policy, avg_salary / divisor_n)
    hourly_rate = _money(policy, daily_rate / hours_day)
    if full_month:
        base = _money(policy, avg_salary)
    else:
        base = _money(policy, daily_rate * Decimal(len(calendar_days)))

    taken = paid_leave
    unused = max(ZERO, allowance - taken)
    unused = max(ZERO, unused - paid_out_days)
    max_cf = _dec(policy.get("max_carry_forward"))
    if _bool(policy.get("leave_carry_forward")) and max_cf > 0:
        unused = min(unused, unused)  # carry handled at finalize
    comp = ZERO
    if unused_ok and payout_now and unused > 0:
        comp = _money(policy, unused * daily_rate)

    allow_total = ZERO
    for item in person.get("allowances") or []:
        if not isinstance(item, dict):
            continue
        allow_total += _dec(item.get("amount"))
    bonus = ZERO
    other_ded = ZERO
    adj_add = ZERO
    adj_ded = ZERO
    for adj in adjustments:
        amt = _dec(adj.get("amount"))
        kind = str(adj.get("kind") or "")
        if kind in {"bonus"}:
            bonus += amt
        elif kind in {"allowance"}:
            allow_total += amt
        elif kind in {"deduction"}:
            other_ded += abs(amt)
        else:
            if amt >= 0:
                adj_add += amt
            else:
                adj_ded += abs(amt)

    late_ded = _money(policy, late_ded)
    early_ded = _money(policy, early_ded)
    absence_ded = _money(policy, absence_ded)
    unpaid_ded = _money(policy, unpaid_ded)
    allow_total = _money(policy, allow_total)
    bonus = _money(policy, bonus)
    other_ded = _money(policy, other_ded)
    adj_add = _money(policy, adj_add)
    adj_ded = _money(policy, adj_ded)
    extra_pay = _money(policy, extra_pay)
    additions = _money(policy, comp + allow_total + bonus + adj_add + extra_pay)
    deductions = _money(policy, late_ded + early_ded + absence_ded + unpaid_ded + other_ded + adj_ded)
    net = _money(policy, base + additions - deductions)

    def line(code: str, label: str, amount: Decimal, kind: str) -> None:
        lines.append({"code": code, "label": label, "amount": str(_money(policy, amount)), "kind": kind})

    line("base", "Contractual base salary", base, "info")
    if comp > 0:
        line("unused_leave", "Unused leave compensation", comp, "add")
    if allow_total:
        line("allowance", "Allowances", allow_total, "add")
    if bonus:
        line("bonus", "Bonuses", bonus, "add")
    if extra_pay:
        line("extra_day", "Extra days", extra_pay, "add")
    if adj_add:
        line("adj_add", "Manual additions", adj_add, "add")
    if late_ded:
        line("late", "Lateness deduction", late_ded, "deduct")
    if early_ded:
        line("early", "Early departure deduction", early_ded, "deduct")
    if absence_ded:
        line("absence", "Absence deduction", absence_ded, "deduct")
    if unpaid_ded:
        line("unpaid_leave", "Unpaid leave deduction", unpaid_ded, "deduct")
    if other_ded or adj_ded:
        line("other", "Other deductions", other_ded + adj_ded, "deduct")
    line("net", "Net payable", net, "net")

    inherited = person.get("inherited") or {}
    return {
        "employee_id": person["employee_id"],
        "name": person.get("name") or person["employee_id"],
        "currency": policy.get("currency") or "EGP",
        "monthly_salary": str(_money(policy, profile_salary)),
        "work_type": str(person.get("work_type") or "full").strip().lower() if str(person.get("work_type") or "").strip().lower() in {"part", "leader"} else "full",
        "base_salary": str(base),
        "extra_days": extra_days,
        "extra_pay": str(extra_pay),
        "late_days": late_days,
        "absent_marks": absent_marks,
        "no_clock_in_days": no_clock_in_days,
        "no_clock_out_days": no_clock_out_days,
        "days_off_used": days_off_used_n,
        "days_off_left": max(0, DAYS_OFF_ALLOWANCE - days_off_used_n),
        "daily_rate": str(daily_rate),
        "hourly_rate": str(hourly_rate),
        "divisor": policy.get("divisor"),
        "divisor_value": str(divisor_n),
        "paid_leave_days": str(paid_leave),
        "unpaid_leave_days": str(unpaid_leave),
        "absence_days": str(_money(policy, absence_days)),
        "partial_absence_days": str(_money(policy, partial_days)),
        "late_minutes": late_minutes,
        "chargeable_late_minutes": chargeable_total,
        "late_deduction": str(late_ded),
        "early_deduction": str(early_ded),
        "absence_deduction": str(absence_ded),
        "unpaid_leave_deduction": str(unpaid_ded),
        "unused_leave_days": str(unused),
        "unused_leave_compensation": str(comp),
        "allowances": str(allow_total),
        "bonuses": str(bonus),
        "adjustments": str(_money(policy, adj_add - adj_ded)),
        "other_deductions": str(other_ded),
        "additions_total": str(additions),
        "deductions_total": str(deductions),
        "net_pay": str(net),
        "leave_allowance": str(allowance),
        "flags": missing_flags,
        "inherited": inherited,
        "lines": lines,
        "explanation": explanation,
    }


def _load_people_maps(start: date, end: date, employee_id: str = "", period_id: int | None = None) -> dict:
    policy = merge_policy(get_policy()["values"])
    people = db.list_employees()
    if employee_id:
        people = [p for p in people if p["employee_id"] == employee_id]
    people = [p for p in people if (p.get("employment_status") or "active").lower() != "inactive"]
    ids = [p["employee_id"] for p in people]
    sessions = db.list_sessions(date_from=start.isoformat(), date_to=end.isoformat(), employee_id=employee_id)
    sessions = [s for s in sessions if s.get("employee_id") in set(ids)]
    with db.get_connection() as conn:
        leave = conn.execute(
            "SELECT * FROM leave_ledger WHERE work_date >= ? AND work_date <= ?" + (" AND employee_id = ?" if employee_id else ""),
            (start.isoformat(), end.isoformat(), employee_id) if employee_id else (start.isoformat(), end.isoformat()),
        ).fetchall()
        adjs = []
        if ids:
            placeholders = ",".join("?" * len(ids))
            if period_id is not None:
                adjs = conn.execute(
                    f"SELECT * FROM payroll_adjustments WHERE employee_id IN ({placeholders}) AND (period_id IS NULL OR period_id = ?)",
                    tuple(ids) + (period_id,),
                ).fetchall()
            else:
                adjs = conn.execute(
                    f"SELECT * FROM payroll_adjustments WHERE employee_id IN ({placeholders}) AND period_id IS NULL",
                    tuple(ids),
                ).fetchall()
        profiles = conn.execute("SELECT * FROM employee_payroll").fetchall()
        hist = conn.execute("SELECT * FROM employee_salary_history ORDER BY effective_date").fetchall()
        payouts = conn.execute("SELECT employee_id, SUM(CAST(days AS REAL)) AS days FROM leave_payouts WHERE year = ? GROUP BY employee_id", (str(start.year),)).fetchall()
        balances = conn.execute("SELECT * FROM leave_balances WHERE year = ?", (str(start.year),)).fetchall()
    prof = {r["employee_id"]: dict(r) for r in profiles}
    hist_map: dict[str, list] = {}
    for row in hist:
        hist_map.setdefault(row["employee_id"], []).append(dict(row))
    leave_map: dict[str, list] = {}
    for row in leave:
        leave_map.setdefault(row["employee_id"], []).append(dict(row))
    adj_map: dict[str, list] = {}
    for row in adjs:
        adj_map.setdefault(row["employee_id"], []).append(dict(row))
    paid_map = {r["employee_id"]: _dec(r["days"]) for r in payouts}
    carry_map = {r["employee_id"]: _dec(r["carried"]) for r in balances}
    sess_map: dict[str, list] = {}
    for s in sessions:
        sess_map.setdefault(s["employee_id"], []).append(s)
    return {
        "policy": policy,
        "people": people,
        "prof": prof,
        "hist_map": hist_map,
        "leave_map": leave_map,
        "adj_map": adj_map,
        "paid_map": paid_map,
        "carry_map": carry_map,
        "sess_map": sess_map,
        "working": _working_days(),
        "holidays": _holidays(policy),
        "shifts": db.list_shifts(),
    }


def _person_bundle(person: dict, maps: dict) -> dict:
    row = maps["prof"].get(person["employee_id"]) or {}
    unused = row.get("unused_leave_eligible")
    if unused is not None:
        unused = bool(unused)
    return {
        **person,
        "monthly_salary": row.get("monthly_salary") or "0",
        "divisor_override": row.get("divisor_override"),
        "leave_allowance_override": row.get("leave_allowance_override"),
        "shift_id": row.get("shift_id_override") if row.get("shift_id_override") is not None else person.get("shift_id"),
        "unused_leave_eligible": unused,
        "allowances": json.loads(row["allowances_json"]) if row.get("allowances_json") else [],
        "inherited": {
            "monthly_salary": not row or _dec(row.get("monthly_salary")) == ZERO,
            "divisor": not row.get("divisor_override"),
            "leave_allowance": not row.get("leave_allowance_override"),
            "shift": row.get("shift_id_override") is None,
            "unused_leave_eligible": unused is None,
        },
    }


def _payout_now(policy: dict, start: date, end: date) -> bool:
    if not _bool(policy.get("unused_leave_compensation")):
        return False
    mode = policy.get("unused_payout") or "end_of_period"
    last = (date(end.year, end.month, 1) + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    if mode == "monthly":
        return True
    if mode == "annual":
        return end.month == 12 and end.day >= 28
    return end == last or (end - start).days >= 27


def preview(start_s: str, end_s: str, employee_id: str = "", period_id: int | None = None) -> dict:
    db.init_db()
    start, end = _parse_date(start_s), _parse_date(end_s)
    if end < start:
        raise PayrollError("End date must be on or after the start date.")
    maps = _load_people_maps(start, end, employee_id, period_id)
    policy = maps["policy"]
    payout = _payout_now(policy, start, end)
    rows = []
    for person in maps["people"]:
        bundled = _person_bundle(person, maps)
        if bundled.get("divisor_override") in DIVISORS:
            use_policy = dict(policy)
            use_policy["divisor"] = bundled["divisor_override"]
        else:
            use_policy = policy
        calc = calculate_employee(
            bundled,
            use_policy,
            start,
            end,
            maps["sess_map"].get(person["employee_id"], []),
            maps["leave_map"].get(person["employee_id"], []),
            maps["adj_map"].get(person["employee_id"], []),
            maps["hist_map"].get(person["employee_id"], []),
            maps["working"],
            maps["holidays"],
            maps["shifts"],
            maps["paid_map"].get(person["employee_id"], ZERO),
            maps["carry_map"].get(person["employee_id"], ZERO),
            payout_now=payout and bool(bundled.get("unused_leave_eligible") if bundled.get("unused_leave_eligible") is not None else _bool(policy.get("unused_leave_compensation"))),
        )
        rows.append(calc)
    return _pack(start, end, policy, rows, status="preview")


def month_board(as_of: str = "") -> dict:
    """Salary and attendance totals for the month, through the given day."""
    db.init_db()
    day = _parse_date(as_of) if (as_of or "").strip() else date.today()
    start, last = _month_span(day)
    end = min(day, last)
    packed = preview(start.isoformat(), end.isoformat())
    future_off: dict[str, int] = {}
    if end < last:
        with db.get_connection() as conn:
            extra = conn.execute(
                """
                SELECT employee_id, COUNT(*) AS n FROM leave_ledger
                WHERE kind = 'day_off' AND work_date > ? AND work_date <= ?
                GROUP BY employee_id
                """,
                (end.isoformat(), last.isoformat()),
            ).fetchall()
        future_off = {row["employee_id"]: int(row["n"]) for row in extra}
    policy = merge_policy(get_policy()["values"])
    staff = []
    for emp in packed["employees"]:
        used = int(emp.get("days_off_used") or 0) + future_off.get(emp["employee_id"], 0)
        left = max(0, DAYS_OFF_ALLOWANCE - used)
        salary = _dec(emp.get("monthly_salary"))
        deducted = (
            _dec(emp.get("late_deduction"))
            + _dec(emp.get("absence_deduction"))
            + _dec(emp.get("unpaid_leave_deduction"))
            + _dec(emp.get("early_deduction"))
        )
        deducted = _money(policy, deducted)
        staff.append(
            {
                "employee_id": emp["employee_id"],
                "name": emp.get("name") or emp["employee_id"],
                "salary": str(_money(policy, salary)),
                "days_off_left": left,
                "days_off_used": min(used, DAYS_OFF_ALLOWANCE),
                "days_off_allowance": DAYS_OFF_ALLOWANCE,
                "late_days": int(emp.get("late_days") or 0),
                "absent": int(emp.get("absent_marks") or 0),
                "no_clock_in": int(emp.get("no_clock_in_days") or 0),
                "no_clock_out": int(emp.get("no_clock_out_days") or 0),
                "extra_days": int(emp.get("extra_days") or 0),
                "extra_pay": str(_money(policy, _dec(emp.get("extra_pay")))),
                "work_type": emp.get("work_type") or "full",
                "deducted": str(deducted),
                "to_pay": str(_money(policy, salary - deducted + _dec(emp.get("extra_pay")))),
            }
        )
    staff.sort(key=lambda row: str(row["name"]).lower())
    return {
        "month": start.strftime("%Y-%m"),
        "through": end.isoformat(),
        "days_off_allowance": DAYS_OFF_ALLOWANCE,
        "currency": packed.get("currency") or "EGP",
        "employees": staff,
    }


def _pack(start: date, end: date, policy: dict, rows: list[dict], status: str, period_id: int | None = None) -> dict:
    total = ZERO
    add = ZERO
    ded = ZERO
    flags = 0
    for row in rows:
        total += _dec(row["net_pay"])
        add += _dec(row["additions_total"])
        ded += _dec(row["deductions_total"])
        if row.get("flags"):
            flags += 1
    return {
        "period_id": period_id,
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
        "status": status,
        "currency": policy.get("currency") or "EGP",
        "employee_count": len(rows),
        "total_payroll": str(_money(policy, total)),
        "total_additions": str(_money(policy, add)),
        "total_deductions": str(_money(policy, ded)),
        "outstanding_review": flags,
        "suspicious": flags,
        "employees": rows,
    }


def calculate_period(start_s: str, end_s: str, actor: str) -> dict:
    db.init_db()
    start, end = _parse_date(start_s), _parse_date(end_s)
    with db.get_connection() as conn:
        existing = conn.execute(
            "SELECT * FROM payroll_periods WHERE start_date = ? AND end_date = ?",
            (start.isoformat(), end.isoformat()),
        ).fetchone()
    if existing and existing["status"] in LOCKED:
        raise PayrollError("This period is finalized. Create an audited correction instead of recalculating.")
    pid_hint = int(existing["id"]) if existing else None
    data = preview(start_s, end_s, period_id=pid_hint)
    policy = get_policy()["values"]
    stamp = _now()
    blob = json.dumps(policy, sort_keys=True)
    with db.get_connection() as conn:
        if existing:
            pid = int(existing["id"])
            conn.execute(
                "UPDATE payroll_periods SET status = 'calculated', policy_snapshot = ?, created_at = ?, created_by = ? WHERE id = ?",
                (blob, stamp, actor, pid),
            )
            conn.execute("DELETE FROM payroll_rows WHERE period_id = ?", (pid,))
        else:
            conn.execute(
                "INSERT INTO payroll_periods (start_date, end_date, status, policy_snapshot, created_at, created_by) VALUES (?, ?, 'calculated', ?, ?, ?)",
                (start.isoformat(), end.isoformat(), blob, stamp, actor),
            )
            pid = int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])
        for row in data["employees"]:
            conn.execute(
                "INSERT INTO payroll_rows (period_id, employee_id, snapshot, net_pay) VALUES (?, ?, ?, ?)",
                (pid, row["employee_id"], json.dumps(row), row["net_pay"]),
            )
    _audit(actor, "period.calculate", "payroll_periods", str(pid), "", data["total_payroll"])
    return get_period(pid)


def _rows_for(period_id: int) -> list[dict]:
    with db.get_connection() as conn:
        rows = conn.execute("SELECT snapshot FROM payroll_rows WHERE period_id = ? ORDER BY employee_id", (period_id,)).fetchall()
    out = []
    for row in rows:
        try:
            out.append(json.loads(row["snapshot"]))
        except json.JSONDecodeError:
            continue
    return out


def get_period(period_id: int) -> dict:
    db.init_db()
    period = _period(period_id)
    policy = merge_policy(json.loads(period["policy_snapshot"]) if period.get("policy_snapshot") else None)
    rows = _rows_for(period_id)
    packed = _pack(_parse_date(period["start_date"]), _parse_date(period["end_date"]), policy, rows, period["status"], period_id)
    packed["approved_by"] = period.get("approved_by") or ""
    packed["finalized_by"] = period.get("finalized_by") or ""
    packed["approved_at"] = period.get("approved_at") or ""
    packed["finalized_at"] = period.get("finalized_at") or ""
    return packed


def list_periods() -> list[dict]:
    db.init_db()
    with db.get_connection() as conn:
        rows = conn.execute("SELECT id, start_date, end_date, status, created_at, created_by FROM payroll_periods ORDER BY start_date DESC").fetchall()
    return [dict(r) for r in rows]


def set_status(period_id: int, status: str, actor: str, role: str) -> dict:
    db.init_db()
    if status not in PERIOD_STATES:
        raise PayrollError("Unknown payroll status.")
    period = _period(period_id)
    if period["status"] in LOCKED:
        raise PayrollError("Finalized payroll is locked.")
    if status == "finalized" and role != "super":
        raise PayrollError("Only the company owner can finalize payroll.")
    if status == "finalized":
        return finalize(period_id, actor)
    stamp = _now()
    fields = {"status": status}
    sql = "UPDATE payroll_periods SET status = ?"
    args: list[Any] = [status]
    if status == "approved":
        sql += ", approved_at = ?, approved_by = ?"
        args.extend([stamp, actor])
    sql += " WHERE id = ?"
    args.append(period_id)
    with db.get_connection() as conn:
        conn.execute(sql, args)
    _audit(actor, "period." + status, "payroll_periods", str(period_id), period["status"], status)
    return get_period(period_id)


def finalize(period_id: int, actor: str) -> dict:
    db.init_db()
    period = _period(period_id)
    if period["status"] in LOCKED:
        return get_period(period_id)
    rows = _rows_for(period_id)
    year = str(_parse_date(period["start_date"]).year)
    policy = merge_policy(json.loads(period["policy_snapshot"]))
    stamp = _now()
    with db.get_connection() as conn:
        try:
            conn.execute(
                "UPDATE payroll_periods SET status = 'finalized', finalized_at = ?, finalized_by = ? WHERE id = ? AND status != 'finalized'",
                (stamp, actor, period_id),
            )
            if conn.execute("SELECT changes()").fetchone()[0] == 0:
                raise PayrollError("Payroll was already finalized.")
            for row in rows:
                unused = _dec(row.get("unused_leave_days"))
                comp = _dec(row.get("unused_leave_compensation"))
                if unused > 0 and comp > 0:
                    try:
                        conn.execute(
                            "INSERT INTO leave_payouts (employee_id, period_id, year, days, amount, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                            (row["employee_id"], period_id, year, str(unused), str(comp), stamp),
                        )
                    except Exception as exc:
                        raise PayrollError("Unused leave for this employee was already paid for this period.") from exc
                    prev = conn.execute(
                        "SELECT paid_out, carried FROM leave_balances WHERE employee_id = ? AND year = ?",
                        (row["employee_id"], year),
                    ).fetchone()
                    paid = _dec(prev["paid_out"] if prev else 0) + unused
                    conn.execute(
                        "INSERT INTO leave_balances (employee_id, year, carried, paid_out) VALUES (?, ?, '0', ?) ON CONFLICT(employee_id, year) DO UPDATE SET paid_out = excluded.paid_out, carried = '0'",
                        (row["employee_id"], year, str(paid)),
                    )
                elif _bool(policy.get("leave_carry_forward")):
                    leftover = unused
                    cap = _dec(policy.get("max_carry_forward"))
                    if cap > 0:
                        leftover = min(leftover, cap)
                    conn.execute(
                        "INSERT INTO leave_balances (employee_id, year, carried, paid_out) VALUES (?, ?, ?, '0') ON CONFLICT(employee_id, year) DO UPDATE SET carried = excluded.carried",
                        (row["employee_id"], year, str(leftover)),
                    )
        except Exception:
            raise
    _audit(actor, "period.finalize", "payroll_periods", str(period_id), period["status"], "finalized")
    return get_period(period_id)


def employee_detail(period_id: int, employee_id: str) -> dict:
    period = get_period(period_id)
    for row in period["employees"]:
        if row["employee_id"] == employee_id:
            return {"period": {k: period[k] for k in ("period_id", "start_date", "end_date", "status", "currency")}, "employee": row}
    raise PayrollError("Employee is not on this payroll.")


def export_period(period_id: int, fmt: str) -> tuple[bytes, str, str]:
    data = get_period(period_id)
    fmt = (fmt or "xlsx").lower()
    headers = [
        "Employee ID",
        "Name",
        "Period start",
        "Period end",
        "Base salary",
        "Paid leave days",
        "Unpaid leave days",
        "Unused leave compensation",
        "Late minutes",
        "Late deduction",
        "Absence deduction",
        "Early deduction",
        "Allowances",
        "Bonuses",
        "Adjustments",
        "Total additions",
        "Total deductions",
        "Net salary",
        "Status",
    ]
    rows = []
    for e in data["employees"]:
        rows.append(
            [
                e["employee_id"],
                e["name"],
                data["start_date"],
                data["end_date"],
                e["base_salary"],
                e["paid_leave_days"],
                e["unpaid_leave_days"],
                e["unused_leave_compensation"],
                e["late_minutes"],
                e["late_deduction"],
                e["absence_deduction"],
                e["early_deduction"],
                e["allowances"],
                e["bonuses"],
                e["adjustments"],
                e["additions_total"],
                e["deductions_total"],
                e["net_pay"],
                data["status"],
            ]
        )
    stamp = f"{data['start_date']}_to_{data['end_date']}"
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(rows)
        payload = ("\ufeff" + buf.getvalue()).encode("utf-8")
        return payload, f"BunkerAttendance_Payroll_{stamp}.csv", "text/csv; charset=utf-8"
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append(["Bunker Attendance payroll", data["status"], data.get("currency") or "EGP"])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 18
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f"BunkerAttendance_Payroll_{stamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def payslip(period_id: int, employee_id: str) -> dict:
    detail = employee_detail(period_id, employee_id)
    emp = detail["employee"]
    return {
        "title": "Payslip",
        "period": detail["period"],
        "employee_id": emp["employee_id"],
        "name": emp["name"],
        "currency": emp.get("currency") or "EGP",
        "lines": emp.get("lines") or [],
        "net_pay": emp["net_pay"],
        "explanation": emp.get("explanation") or [],
        "status": detail["period"]["status"],
    }

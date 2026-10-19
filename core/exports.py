"""Absence Less attendance/employee exports. Read-only: never writes attendance rows."""

from __future__ import annotations

import csv
import io
import os
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core import db
from core.reports import attach_salary
from core.time_utils import hours_between

MAX_DAYS = 366
MAX_ROWS = 50_000
REPORT_TYPES = ("daily", "monthly", "employees", "late", "absences", "hours")
FORMATS = ("xlsx", "csv", "pdf")

DETAIL_DEFAULT = [
    "employee_id",
    "name",
    "date",
    "check_in",
    "check_out",
    "hours",
    "status",
]
DETAIL_ALL = DETAIL_DEFAULT + ["shift", "department", "branch", "late_minutes", "device"]
DIRECTORY_DEFAULT = ["employee_id", "name", "email", "department", "branch", "shift", "face", "status"]
DIRECTORY_ALL = DIRECTORY_DEFAULT + ["phone", "job_title", "joining_date"]
SUMMARY_HEADERS = ["Employee", "Days present", "Late arrivals", "Absences", "Total hours", "Base salary", "Net pay"]

INK = PatternFill("solid", fgColor="001F3F")
CREAM = PatternFill("solid", fgColor="FFFFFF")
GREEN = PatternFill("solid", fgColor="C6F4D6")
LATE = PatternFill("solid", fgColor="FCE4C4")
RED = PatternFill("solid", fgColor="FFC7C7")
PURPLE = PatternFill("solid", fgColor="E3D0F2")
ZEBRA = PatternFill("solid", fgColor="F4F5EE")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Calibri")
TITLE_FONT = Font(bold=True, color="001F3F", size=16, name="Calibri")
THIN = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
PDF_INK = colors.HexColor("#001F3F")
PDF_CREAM = colors.HexColor("#FFFFFF")
PDF_GREEN = colors.HexColor("#C6F4D6")
PDF_LATE = colors.HexColor("#FCE4C4")
PDF_RED = colors.HexColor("#FFC7C7")
PDF_PURPLE = colors.HexColor("#E3D0F2")

_DAY_INDEX = {"mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6}
_FONT_REGISTERED = False


class ExportError(ValueError):
    pass


def _dec(value: float | None) -> Decimal | None:
    if value is None:
        return None
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _hours_out(value: float | None) -> float | None:
    d = _dec(value)
    return float(d) if d is not None else None


def csv_safe(value) -> str:
    text = "" if value is None else str(value)
    if text[:1] in {"=", "+", "-", "@", "\t", "\r"}:
        return "'" + text
    return text


def _setting(key: str, default: str = "") -> str:
    try:
        with db.get_connection() as conn:
            row = conn.execute("SELECT value FROM gym_settings WHERE key = ?", (key,)).fetchone()
        if row and row[0] not in (None, ""):
            return str(row[0])
    except Exception:
        pass
    return default


def _status_wanted(label: str, status_filter: str) -> bool:
    if not (status_filter or "").strip():
        return True
    want = status_filter.strip().lower()
    if want in {"missing check-out", "missing checkout", "incomplete"}:
        return label.lower() == "incomplete"
    return label.lower() == want


def _elapsed_hours(session: dict) -> float | None:
    if not session.get("clock_out"):
        return None
    if session.get("hours") is not None:
        return _hours_out(float(session["hours"]))
    try:
        start = datetime.fromisoformat(str(session["clock_in"]))
        end = datetime.fromisoformat(str(session["clock_out"]))
    except (TypeError, ValueError):
        return None
    return _hours_out(hours_between(start, end))


def parse_working_days(text: str) -> set[int]:
    raw = (text or "Sun-Thu").strip().lower()
    if "-" in raw and "," not in raw:
        left, right = raw.split("-", 1)
        a = _DAY_INDEX.get(left.strip()[:3])
        b = _DAY_INDEX.get(right.strip()[:3])
        if a is not None and b is not None:
            out = set()
            k = a
            while True:
                out.add(k)
                if k == b:
                    break
                k = (k + 1) % 7
            return out
    found = {_DAY_INDEX[p.strip()[:3]] for p in raw.replace(";", ",").split(",") if p.strip()[:3] in _DAY_INDEX}
    return found or {6, 0, 1, 2, 3}


def _dates_between(start: date, end: date) -> list[date]:
    days = []
    cur = start
    while cur <= end:
        days.append(cur)
        cur += timedelta(days=1)
    return days


def _parse_day(value: str, fallback: date) -> date:
    text = (value or "").strip()
    if not text:
        return fallback
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ExportError("Dates must be YYYY-MM-DD.") from exc


def validate_range(start_date: str, end_date: str) -> tuple[date, date]:
    today = date.today()
    start = _parse_day(start_date, today)
    end = _parse_day(end_date, today)
    if end < start:
        raise ExportError("End date must be on or after the start date.")
    if (end - start).days + 1 > MAX_DAYS:
        raise ExportError(f"Date range cannot exceed {MAX_DAYS} days.")
    return start, end


def _fmt_date(iso: str) -> str:
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(iso or "")


def _fmt_time(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(str(iso)).strftime("%H:%M")
    except ValueError:
        return str(iso)


def _filter_people(
    people: Iterable[dict],
    *,
    branch: str = "",
    department: str = "",
    employee_id: str = "",
    include_inactive: bool = False,
) -> list[dict]:
    out = []
    for person in people:
        status = (person.get("employment_status") or "active").lower()
        if not include_inactive and status == "inactive":
            continue
        if employee_id and person.get("employee_id") != employee_id:
            continue
        if branch and (person.get("branch") or "HQ") != branch:
            continue
        if department and (person.get("department") or "") != department:
            continue
        out.append(person)
    return out


def collect(params: dict) -> dict:
    report = (params.get("report") or "monthly").strip().lower()
    if report not in REPORT_TYPES:
        raise ExportError("Unknown report type.")
    start, end = validate_range(params.get("start_date") or "", params.get("end_date") or "")
    branch = (params.get("branch") or params.get("branch_id") or "").strip()
    department = (params.get("department") or params.get("department_id") or "").strip()
    employee_id = (params.get("employee_id") or "").strip()
    status_filter = (params.get("status") or "").strip()
    include_inactive = str(params.get("include_inactive") or "") in {"1", "true", "yes"}
    grace = 0
    try:
        grace = max(0, int(_setting("grace_minutes", "10") or 0))
    except ValueError:
        grace = 10
    working = parse_working_days(_setting("working_days", "Sun-Thu"))
    people = _filter_people(
        db.list_employees(),
        branch=branch,
        department=department,
        employee_id=employee_id,
        include_inactive=include_inactive,
    )
    allowed = {p["employee_id"] for p in people}
    by_id = {p["employee_id"]: p for p in people}
    sessions = db.list_sessions(date_from=start.isoformat(), date_to=end.isoformat(), employee_id=employee_id)
    sessions = [s for s in sessions if s.get("employee_id") in allowed]
    if len(sessions) > MAX_ROWS:
        raise ExportError("Export is too large. Narrow the date range or filters.")

    grouped: dict[tuple[str, str], list[dict]] = {}
    for session in sessions:
        eid = str(session.get("employee_id") or "")
        work_day = str(session.get("work_date") or "")
        grouped.setdefault((eid, work_day), []).append(session)

    details = []
    present_days: dict[str, set[str]] = {eid: set() for eid in allowed}
    late_days: dict[str, set[str]] = {eid: set() for eid in allowed}
    hours_sum: dict[str, Decimal] = {eid: Decimal("0.00") for eid in allowed}

    for (eid, work_day), group in grouped.items():
        person = by_id.get(eid) or {}
        group = sorted(group, key=lambda s: str(s.get("clock_in") or ""))
        closed = [s for s in group if s.get("clock_out")]
        open_rows = [s for s in group if not s.get("clock_out")]
        hour_parts = [_elapsed_hours(s) for s in closed]
        hour_parts = [h for h in hour_parts if h is not None]
        hours = _hours_out(sum(hour_parts)) if hour_parts else None
        late_min = max(int(s.get("minutes_late") or 0) for s in group)
        if open_rows:
            label = "Incomplete"
        elif late_min > grace:
            label = "Late"
        else:
            label = "Present"
        if label in {"Present", "Late", "Incomplete"}:
            try:
                work_dt = date.fromisoformat(work_day)
            except ValueError:
                work_dt = None
            if work_dt is None or work_dt.weekday() in working:
                present_days[eid].add(work_day)
        if label == "Late":
            late_days[eid].add(work_day)
        if hours is not None:
            hours_sum[eid] += Decimal(str(hours))
        last_out = ""
        if closed and not open_rows:
            last = max(closed, key=lambda s: str(s.get("clock_out") or ""))
            last_out = _fmt_time(last.get("clock_out"))
        row = {
            "employee_id": eid,
            "name": group[0].get("employee_name") or person.get("name") or "",
            "date": _fmt_date(work_day),
            "date_iso": work_day,
            "check_in": _fmt_time(group[0].get("clock_in")),
            "check_out": last_out,
            "hours": hours,
            "status": label,
            "shift": group[0].get("shift_name") or person.get("shift_name") or "",
            "department": person.get("department") or "",
            "branch": person.get("branch") or "HQ",
            "late_minutes": late_min,
            "device": group[0].get("device_mac") or "",
        }
        if report == "late" and label != "Late":
            continue
        if report == "hours" and hours is None:
            continue
        if not _status_wanted(label, status_filter):
            continue
        details.append(row)

    details.sort(key=lambda r: (r.get("date_iso") or "", (r["name"] or "").lower(), r["employee_id"]))

    scheduled = [d for d in _dates_between(start, end) if d.weekday() in working]
    absences = []
    summary = []
    for person in sorted(people, key=lambda p: (p.get("name") or "").lower()):
        eid = person["employee_id"]
        present = present_days.get(eid, set())
        absent_n = 0
        for day in scheduled:
            key = day.isoformat()
            if key not in present:
                absent_n += 1
                if report in {"absences", "monthly", "daily"}:
                    absences.append(
                        {
                            "employee_id": eid,
                            "name": person.get("name") or "",
                            "date": _fmt_date(key),
                            "date_iso": key,
                            "check_in": "",
                            "check_out": "",
                            "hours": None,
                            "status": "Absent",
                            "shift": person.get("shift_name") or "",
                            "department": person.get("department") or "",
                            "branch": person.get("branch") or "HQ",
                            "late_minutes": 0,
                            "device": "",
                        }
                    )
        total = hours_sum.get(eid, Decimal("0.00"))
        summary.append(
            {
                "employee": person.get("name") or eid,
                "employee_id": eid,
                "days_present": len(present),
                "late_arrivals": len(late_days.get(eid, set())),
                "absences": absent_n,
                "total_hours": float(total),
            }
        )
    attach_salary(summary, start.isoformat(), end.isoformat(), employee_id=employee_id)

    if report == "absences":
        if status_filter and status_filter.lower() != "absent":
            details = []
        else:
            details = sorted(absences, key=lambda r: (r.get("date_iso") or "", (r["name"] or "").lower()))

    directory = []
    for person in sorted(people, key=lambda p: (p.get("name") or "").lower()):
        directory.append(
            {
                "employee_id": person.get("employee_id"),
                "name": person.get("name"),
                "email": person.get("email") or "",
                "phone": person.get("phone") or "",
                "job_title": person.get("job_title") or "",
                "department": person.get("department") or "",
                "branch": person.get("branch") or "HQ",
                "shift": person.get("shift_name") or "",
                "face": "Enrolled" if person.get("has_face") else "Not enrolled",
                "status": person.get("employment_status") or "active",
                "joining_date": person.get("joining_date") or "",
            }
        )

    return {
        "report": report,
        "start": start,
        "end": end,
        "details": details,
        "summary": summary,
        "directory": directory,
        "grace": grace,
        "company": _setting("company_name", "Bunker Attendance") or "Bunker Attendance",
        "working_days": _setting("working_days", "Sun-Thu"),
        "hours_note": "Hours are elapsed check-in to check-out. Salary is calculated from payroll settings for this range.",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "branch": branch or "All branches",
        "department": department or "All departments",
        "count": len(directory) if report == "employees" else len(details),
    }


def selected_columns(report: str, raw: str) -> list[str]:
    allowed = DIRECTORY_ALL if report == "employees" else DETAIL_ALL
    defaults = DIRECTORY_DEFAULT if report == "employees" else DETAIL_DEFAULT
    if not (raw or "").strip():
        return list(defaults)
    picked = []
    for key in raw.split(","):
        key = key.strip()
        if key in allowed and key not in picked:
            picked.append(key)
    return picked or list(defaults)


def column_labels(keys: list[str]) -> list[str]:
    labels = {
        "employee_id": "Employee ID",
        "name": "Name",
        "date": "Date",
        "check_in": "Check-in",
        "check_out": "Check-out",
        "hours": "Hours",
        "status": "Status",
        "shift": "Shift",
        "department": "Department",
        "branch": "Branch",
        "late_minutes": "Late minutes",
        "device": "Device",
        "email": "Email",
        "phone": "Phone",
        "job_title": "Job title",
        "face": "Face enrollment",
        "joining_date": "Joining date",
    }
    return [labels.get(k, k) for k in keys]


def _detail_value(row: dict, key: str):
    value = row.get(key)
    if key == "hours":
        return None if value in (None, "") else float(value)
    return csv_safe(value) if isinstance(value, str) else value


def filename_for(report: str, fmt: str, start: date, end: date) -> str:
    if report == "employees":
        stamp = date.today().isoformat()
        return f"BunkerAttendance_Employees_{stamp}.{fmt}"
    if report == "monthly" and start.month == end.month and start.year == end.year:
        stamp = start.strftime("%Y-%m")
    else:
        stamp = f"{start.isoformat()}_to_{end.isoformat()}"
    names = {
        "daily": "Attendance",
        "monthly": "Attendance",
        "late": "Late_Arrivals",
        "absences": "Absences",
        "hours": "Working_Hours",
    }
    prefix = names.get(report, "Attendance")
    if fmt == "pdf" and report in {"monthly", "daily", "hours"}:
        return f"BunkerAttendance_Attendance_Summary_{stamp}.pdf"
    return f"BunkerAttendance_{prefix}_{stamp}.{fmt}"


def media_type(fmt: str) -> str:
    if fmt == "xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "pdf":
        return "application/pdf"
    return "text/csv; charset=utf-8"


def build(params: dict, actor: str) -> tuple[bytes, str, str]:
    fmt = (params.get("format") or "xlsx").strip().lower()
    report = (params.get("report") or "monthly").strip().lower()
    date_from = str(params.get("start_date") or "")
    date_to = str(params.get("end_date") or "")
    scope = f"branch={params.get('branch') or params.get('branch_id') or ''};dept={params.get('department') or ''}"
    try:
        if fmt not in FORMATS:
            raise ExportError("Format must be xlsx, csv, or pdf.")
        data = collect(params)
        date_from = data["start"].isoformat()
        date_to = data["end"].isoformat()
        report = data["report"]
        scope = f"branch={data['branch']};dept={data['department']};rows={data['count']}"
        cols = selected_columns(report, params.get("columns") or "")
        if fmt == "xlsx":
            payload = _excel(data, cols)
        elif fmt == "csv":
            payload = _csv(data, cols)
        else:
            payload = _pdf(data, cols)
        db.log_export(
            actor=actor,
            report_type=report,
            fmt=fmt,
            date_from=date_from,
            date_to=date_to,
            scope=scope,
            outcome="ok",
        )
        name = filename_for(report, fmt, data["start"], data["end"])
        return payload, name, media_type(fmt)
    except ExportError:
        db.log_export(
            actor=actor,
            report_type=report,
            fmt=fmt,
            date_from=date_from,
            date_to=date_to,
            scope=scope[:500],
            outcome="rejected",
        )
        raise
    except Exception:
        db.log_export(
            actor=actor,
            report_type=report,
            fmt=fmt,
            date_from=date_from,
            date_to=date_to,
            scope=scope[:500],
            outcome="error",
        )
        raise


def _banner(ws, data: dict, title: str, width: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"] = "Bunker Attendance"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = title
    ws["A2"].font = Font(size=12, color="333333", name="Calibri")
    ws["A3"] = "Period"
    ws["B3"] = f"{data['start'].isoformat()} to {data['end'].isoformat()}"
    ws["A4"] = "Generated"
    ws["B4"] = data["generated_at"].replace("T", " ")
    ws["A5"] = "Company / branch"
    ws["B5"] = f"{data['company']} · {data['branch']}"
    ws["A6"] = "Hours note"
    ws.merge_cells(start_row=6, start_column=2, end_row=6, end_column=width)
    ws["B6"] = data["hours_note"]
    return 8


def _style_header(ws, row: int, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row, col)
        cell.fill = INK
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(width)}{max(ws.max_row, row)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = f"{row}:{row}"
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _status_fills(ws, header_row: int, status_col: int, last_row: int) -> None:
    if last_row <= header_row:
        return
    col = get_column_letter(status_col)
    rng = f"{col}{header_row + 1}:{col}{last_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Present"'], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Late"'], fill=LATE))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Absent"'], fill=RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"On Leave"'], fill=PURPLE))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Incomplete"'], fill=PURPLE))


def _excel(data: dict, cols: list[str]) -> bytes:
    wb = Workbook()
    report = data["report"]
    if report == "employees":
        ws = wb.active
        ws.title = "Employees"
        labels = column_labels(cols)
        header_row = _banner(ws, data, "Employee Directory", max(2, len(cols)))
        for idx, label in enumerate(labels, start=1):
            ws.cell(header_row, idx, label)
        for row in data["directory"]:
            ws.append([csv_safe(row.get(k) or "") for k in cols])
        last = ws.max_row
        for r in range(header_row + 1, last + 1):
            for c in range(1, len(cols) + 1):
                cell = ws.cell(r, c)
                cell.border = THIN
                if r % 2 == 0:
                    cell.fill = ZEBRA
        _style_header(ws, header_row, len(cols))
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
    else:
        ws = wb.active
        ws.title = "Attendance Details"
        labels = column_labels(cols)
        header_row = _banner(ws, data, _title(report), max(2, len(cols)))
        for idx, label in enumerate(labels, start=1):
            ws.cell(header_row, idx, label)
        for item in data["details"]:
            excel_row = []
            for key in cols:
                val = _detail_value(item, key)
                if key in {"name", "employee_id", "status", "shift", "department", "branch", "device"}:
                    val = csv_safe(val)
                excel_row.append(val)
            ws.append(excel_row)
        last = ws.max_row
        hours_idx = cols.index("hours") + 1 if "hours" in cols else None
        status_idx = cols.index("status") + 1 if "status" in cols else None
        date_idx = cols.index("date") + 1 if "date" in cols else None
        for r in range(header_row + 1, last + 1):
            for c in range(1, len(cols) + 1):
                cell = ws.cell(r, c)
                cell.border = THIN
                if r % 2 == 0:
                    cell.fill = ZEBRA
            if hours_idx:
                cell = ws.cell(r, hours_idx)
                if cell.value in ("", None):
                    cell.value = None
                else:
                    cell.number_format = "0.00"
            if date_idx:
                ws.cell(r, date_idx).alignment = Alignment(horizontal="center")
        _style_header(ws, header_row, max(1, len(cols)))
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{max(last, header_row)}"
        if status_idx:
            _status_fills(ws, header_row, status_idx, last)
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 16
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 22

        if report in {"monthly", "daily", "hours"}:
            sm = wb.create_sheet("Employee Summary")
            sh = _banner(sm, data, "Employee Summary — hours and calculated salary", 7)
            for idx, label in enumerate(SUMMARY_HEADERS, start=1):
                sm.cell(sh, idx, label)
            total_present = total_late = total_abs = Decimal("0")
            total_hours = Decimal("0.00")
            total_base = Decimal("0.00")
            total_net = Decimal("0.00")
            for item in data["summary"]:
                sm.append(
                    [
                        csv_safe(item["employee"]),
                        item["days_present"],
                        item["late_arrivals"],
                        item["absences"],
                        item["total_hours"],
                        float(item.get("base_salary") or 0),
                        float(item.get("net_pay") or 0),
                    ]
                )
                total_present += item["days_present"]
                total_late += item["late_arrivals"]
                total_abs += item["absences"]
                total_hours += Decimal(str(item["total_hours"]))
                total_base += Decimal(str(item.get("base_salary") or 0))
                total_net += Decimal(str(item.get("net_pay") or 0))
            sm.append(
                ["Total", int(total_present), int(total_late), int(total_abs), float(total_hours), float(total_base), float(total_net)]
            )
            last_s = sm.max_row
            for r in range(sh, last_s + 1):
                for c in range(1, 8):
                    cell = sm.cell(r, c)
                    cell.border = THIN
                    if r > sh and r % 2 == 0 and r < last_s:
                        cell.fill = ZEBRA
                if r > sh:
                    sm.cell(r, 5).number_format = "0.00"
                    sm.cell(r, 6).number_format = "0.00"
                    sm.cell(r, 7).number_format = "0.00"
            for c in range(1, 8):
                sm.cell(last_s, c).font = Font(bold=True)
                sm.cell(last_s, c).fill = CREAM
            _style_header(sm, sh, 7)
            sm.auto_filter.ref = f"A{sh}:G{max(last_s - 1, sh)}"
            sm.column_dimensions["A"].width = 28
            for col in "BCDEFG":
                sm.column_dimensions[col].width = 16
            sm["A6"] = "Hours note"
            sm["B6"] = data["hours_note"]

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _title(report: str) -> str:
    return {
        "daily": "Daily Attendance",
        "monthly": "Monthly Attendance",
        "late": "Late Arrivals",
        "absences": "Absences",
        "hours": "Working Hours (elapsed time)",
        "employees": "Employee Directory",
    }.get(report, "Attendance")


def _csv(data: dict, cols: list[str]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    if data["report"] == "employees":
        writer.writerow(column_labels(cols))
        for row in data["directory"]:
            writer.writerow([csv_safe(row.get(k) or "") for k in cols])
    else:
        writer.writerow(column_labels(cols))
        for row in data["details"]:
            out = []
            for key in cols:
                val = _detail_value(row, key)
                if val is None:
                    out.append("")
                elif key == "hours":
                    out.append(f"{float(val):.2f}")
                else:
                    out.append(csv_safe(val) if not isinstance(val, (int, float)) else val)
            writer.writerow(out)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _register_pdf_font() -> str:
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return "CODAISans"
    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\tahoma.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if os.path.isfile(path):
            try:
                pdfmetrics.registerFont(TTFont("CODAISans", path))
                _FONT_REGISTERED = True
                return "CODAISans"
            except Exception:
                continue
    _FONT_REGISTERED = True
    return "Helvetica"


def _shape(text: str) -> str:
    value = csv_safe("" if text is None else str(text))
    if any("\u0600" <= ch <= "\u06FF" for ch in value):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display

            return get_display(arabic_reshaper.reshape(value))
        except Exception:
            return value
    return value


def _pdf(data: dict, cols: list[str]) -> bytes:
    font = _register_pdf_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        title=f"Bunker Attendance {_title(data['report'])}",
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=32,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("codaiTitle", parent=styles["Title"], fontName=font, textColor=PDF_INK, fontSize=18)
    body = ParagraphStyle("codaiBody", parent=styles["Normal"], fontName=font, fontSize=9, leading=12)
    heading = ParagraphStyle("codaiH", parent=styles["Heading2"], fontName=font, textColor=PDF_INK, fontSize=12)

    def cell(value) -> Paragraph:
        return Paragraph(_shape("" if value is None else str(value)).replace("&", "&amp;").replace("<", "&lt;"), body)

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(PDF_INK)
        canvas.setFont(font, 8)
        canvas.drawString(24, 16, f"Bunker Attendance · {data['generated_at'].replace('T', ' ')}")
        canvas.drawRightString(landscape(A4)[0] - 24, 16, f"Page {doc_.page}")
        canvas.restoreState()

    story = [
        Paragraph("Bunker Attendance", title),
        Paragraph(_shape(_title(data["report"])), heading),
        Paragraph(
            _shape(
                f"{data['company']} · {data['branch']} · {data['start']} to {data['end']} · {data['hours_note']}"
            ),
            body,
        ),
        Spacer(1, 10),
    ]
    if data["report"] != "employees" and data["summary"]:
        present = sum(x["days_present"] for x in data["summary"])
        late = sum(x["late_arrivals"] for x in data["summary"])
        absn = sum(x["absences"] for x in data["summary"])
        hours = sum(x["total_hours"] for x in data["summary"])
        story.append(
            Paragraph(
                f"People {len(data['summary'])} · Present days {present} · Late {late} · Absences {absn} · "
                f"Elapsed hours {hours:.2f}",
                body,
            )
        )
        story.append(Spacer(1, 8))

    if data["report"] == "employees":
        table_data = [column_labels(cols)]
        for row in data["directory"]:
            table_data.append([cell(row.get(k) or "") for k in cols])
    else:
        table_data = [column_labels(cols)]
        for row in data["details"]:
            cells = []
            for key in cols:
                val = _detail_value(row, key)
                if key == "hours":
                    cells.append("" if val in ("", None) else f"{float(val):.2f}")
                else:
                    cells.append(cell(val))
            table_data.append(cells)

    table = Table(table_data, repeatRows=1)
    cmds = [
        ("FONTNAME", (0, 0), (-1, 0), font),
        ("BACKGROUND", (0, 0), (-1, 0), PDF_INK),
        ("TEXTCOLOR", (0, 0), (-1, 0), PDF_CREAM),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ]
    status_idx = cols.index("status") if "status" in cols else None
    if status_idx is not None:
        rows_src = data["directory"] if data["report"] == "employees" else data["details"]
        fills = {"Present": PDF_GREEN, "Late": PDF_LATE, "Absent": PDF_RED, "Incomplete": PDF_PURPLE, "On Leave": PDF_PURPLE}
        for i, row in enumerate(rows_src, start=1):
            fill = fills.get(str(row.get("status") or ""))
            if fill:
                cmds.append(("BACKGROUND", (status_idx, i), (status_idx, i), fill))
    table.setStyle(TableStyle(cmds))
    story.append(table)

    if data["report"] in {"monthly", "daily", "hours"} and data["summary"]:
        story.append(Spacer(1, 14))
        story.append(Paragraph("Employee Summary", heading))
        story.append(Paragraph(_shape(data["hours_note"]), body))
        sm = [SUMMARY_HEADERS]
        for item in data["summary"]:
            sm.append(
                [
                    cell(item["employee"]),
                    str(item["days_present"]),
                    str(item["late_arrivals"]),
                    str(item["absences"]),
                    f"{item['total_hours']:.2f}",
                    str(item.get("base_salary") or "0.00"),
                    str(item.get("net_pay") or "0.00"),
                ]
            )
        st = Table(sm, repeatRows=1)
        st.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), font),
                    ("BACKGROUND", (0, 0), (-1, 0), PDF_INK),
                    ("TEXTCOLOR", (0, 0), (-1, 0), PDF_CREAM),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(st)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def meta() -> dict:
    people = db.list_employees()
    branches = sorted({(p.get("branch") or "HQ") for p in people})
    depts = sorted({p.get("department") for p in people if p.get("department")})
    return {
        "reports": list(REPORT_TYPES),
        "formats": list(FORMATS),
        "detail_columns": DETAIL_ALL,
        "detail_default": DETAIL_DEFAULT,
        "directory_columns": DIRECTORY_ALL,
        "directory_default": DIRECTORY_DEFAULT,
        "branches": branches,
        "departments": depts,
        "employees": [{"employee_id": p["employee_id"], "name": p["name"]} for p in people if (p.get("employment_status") or "active") != "inactive"],
        "max_days": MAX_DAYS,
        "hours_note": "Elapsed check-in to check-out time. Not payroll hours.",
    }
